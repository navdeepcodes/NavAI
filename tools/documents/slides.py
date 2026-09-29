"""Making a slide deck: a presentation for class, a project pitch, a study deck.

Mike gives the deck as a list of slides (title, bullets, image, table, quote,
two columns, speaker notes) and this builds a real .pptx: proper title
placeholders (so the outline and the slide sorter show the titles), bullets
that are real bullets, speaker notes, slide numbers. Two looks: graphite (the
black-and-grey of Mike itself) and light. Text is sized to what a slide holds
-- fewer words, bigger type -- since a slide that can't be read from the back
of a room isn't finished.
"""
from __future__ import annotations

from pathlib import Path

from logs.logger import logger

THEMES = {
    "graphite": dict(bg="0F0F10", title="F3F3F4", body="C9CBD1", muted="84858B",
                     accent="C7C9CF", panel="1B1B1D", line="2C2C30", header="26262A"),
    "light": dict(bg="F7F7F8", title="0E0E10", body="3A3C42", muted="787A80",
                  accent="3B3E45", panel="EEEEF0", line="D6D6DB", header="E2E2E6"),
}
FONT = "Calibri"
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.8


def _rgb(hexstr: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexstr.upper())


def _bg(slide, colors: dict) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(colors["bg"])


def _style_run(run, size: float, color: str, bold: bool = False, italic: bool = False) -> None:
    from pptx.util import Pt
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)


def _text(slide, x, y, w, h, text: str, size: float, color: str, bold=False, italic=False,
          align=None, anchor=None):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.vertical_anchor = anchor or MSO_ANCHOR.TOP
    para = frame.paragraphs[0]
    para.alignment = align or PP_ALIGN.LEFT
    run = para.add_run()
    run.text = text
    _style_run(run, size, color, bold, italic)
    return box


def _bullet_size(items: list) -> float:
    total = sum(len(str(i if isinstance(i, str) else i[0])) for i in items)
    n = len(items)
    if total <= 220 and n <= 5:
        return 26
    if total <= 420 and n <= 6:
        return 22
    if total <= 700:
        return 19
    return 16


def _add_bullets(slide, x, y, w, h, items: list, colors: dict, size: float | None = None) -> None:
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
    from lxml import etree

    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    size = size or _bullet_size(items)
    first = True
    flat: list[tuple[int, str]] = []
    for item in items:
        if isinstance(item, (list, tuple)):
            flat.extend((1, str(sub)) for sub in item)
        else:
            flat.append((0, str(item)))
    for level, text in flat:
        para = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        run = para.add_run()
        run.text = text
        _style_run(run, size - (3 if level else 0), colors["body"] if level else colors["title"])
        para.space_after = Pt(size * 0.55)
        para.line_spacing = 1.08
        pPr = para._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.34 + 0.4 * level))))
        pPr.set("indent", str(-int(Inches(0.3))))
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", colors["accent"].upper())
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if not level else "–")


def _title(slide, text: str, colors: dict, size: float = 34) -> None:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches
    shape = slide.shapes.title
    shape.left, shape.top = Inches(MARGIN), Inches(0.55)
    shape.width, shape.height = Inches(SLIDE_W - 2 * MARGIN), Inches(1.1)
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    frame.margin_left = frame.margin_right = 0
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = text
    _style_run(run, size if len(text) < 55 else size - 6, colors["title"], bold=True)
    rule = slide.shapes.add_shape(1, Inches(MARGIN), Inches(1.75), Inches(0.9), Inches(0.05))
    rule.fill.solid()
    rule.fill.fore_color.rgb = _rgb(colors["accent"])
    rule.line.fill.background()


def _fit_picture(slide, path: str, x, y, w, h):
    from PIL import Image
    from pptx.util import Emu
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(path, x + int((w - pw) / 2), y + int((h - ph) / 2), pw, ph)


def _footer(slide, number: int, deck_title: str, colors: dict) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches
    _text(slide, Inches(MARGIN), Inches(7.0), Inches(8), Inches(0.3), deck_title[:70], 11, colors["muted"])
    _text(slide, Inches(SLIDE_W - MARGIN - 1), Inches(7.0), Inches(1), Inches(0.3), str(number), 11,
          colors["muted"], align=PP_ALIGN.RIGHT)


def _table(slide, rows: list, x, y, w, h, colors: dict) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    n_rows, n_cols = len(rows), max(len(r) for r in rows)
    shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, min(h, int(n_rows * 460000)))
    table = shape.table
    size = 18 if n_rows <= 5 and n_cols <= 4 else 14 if n_rows <= 9 else 11
    for r, row in enumerate(rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(colors["header"] if r == 0 else colors["panel"])
            frame = cell.text_frame
            frame.word_wrap = True
            para = frame.paragraphs[0]
            run = para.add_run()
            run.text = str(row[c]) if c < len(row) else ""
            _style_run(run, size, colors["title"] if r == 0 else colors["body"], bold=(r == 0))
            para.alignment = PP_ALIGN.LEFT


def _kind(spec: dict, index: int) -> str:
    if spec.get("type"):
        return str(spec["type"]).lower()
    if index == 0 and not any(spec.get(k) for k in ("bullets", "image", "table", "body", "left")):
        return "title"
    if spec.get("table"):
        return "table"
    if spec.get("quote"):
        return "quote"
    if spec.get("left") or spec.get("right"):
        return "two_column"
    if spec.get("image"):
        return "image"
    if not spec.get("bullets") and not spec.get("body") and spec.get("title"):
        return "section"
    return "bullets"


def create_deck(path: str, title: str, slides: list, theme: str = "graphite", author: str = "",
                overwrite: bool = False) -> dict:
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    from tools.filesystem.path_utils import resolve_path

    file = resolve_path(path)
    if file.suffix.lower() != ".pptx":
        return {"status": "error", "retry_safe": True, "error": "A presentation is a .pptx file: name it that way."}
    slides = [s for s in (slides or []) if isinstance(s, dict)]
    if not slides:
        return {"status": "error", "retry_safe": True, "error": "There are no slides to make."}
    existed = file.exists()
    if existed and not overwrite:
        return {"status": "error", "retry_safe": True,
                "error": f"{file.name} already exists in {file.parent}. Nothing was changed. Use a new name, "
                         "or set overwrite to true if the student wants it replaced."}
    if existed:
        try:
            from brain import revert_store
            revert_store.capture(str(file))
        except Exception:
            logger.debug("Couldn't keep a copy before replacing %s", file, exc_info=True)
    colors = THEMES.get((theme or "graphite").lower(), THEMES["graphite"])
    deck_title = title or str(slides[0].get("title") or file.stem)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    prs.core_properties.title = deck_title
    if author:
        prs.core_properties.author = author
    prs.core_properties.comments = "Made by Mike"
    problems: list[str] = []

    for index, spec in enumerate(slides):
        kind = _kind(spec, index)
        heading = str(spec.get("title") or "")
        if kind == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            _bg(slide, colors)
            slide.shapes.title.left, slide.shapes.title.top = Inches(MARGIN + 0.2), Inches(2.4)
            slide.shapes.title.width, slide.shapes.title.height = Inches(SLIDE_W - 2 * MARGIN - 0.4), Inches(1.7)
            tf = slide.shapes.title.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.BOTTOM
            tf.margin_left = 0
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT
            run = para.add_run()
            run.text = heading or deck_title
            _style_run(run, 48 if len(run.text) < 40 else 38, colors["title"], bold=True)
            sub = slide.placeholders[1]
            subtitle = str(spec.get("subtitle") or author or "")
            if subtitle:
                sub.left, sub.top = Inches(MARGIN + 0.2), Inches(4.35)
                sub.width, sub.height = Inches(SLIDE_W - 2 * MARGIN - 0.4), Inches(1.0)
                sub.text_frame.margin_left = 0
                p2 = sub.text_frame.paragraphs[0]
                p2.alignment = PP_ALIGN.LEFT
                r2 = p2.add_run()
                r2.text = subtitle
                _style_run(r2, 22, colors["muted"])
            else:
                sub._element.getparent().remove(sub._element)
            rule = slide.shapes.add_shape(1, Inches(MARGIN + 0.2), Inches(4.2), Inches(1.2), Inches(0.06))
            rule.fill.solid()
            rule.fill.fore_color.rgb = _rgb(colors["accent"])
            rule.line.fill.background()
        elif kind == "section":
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            _bg(slide, colors)
            shape = slide.shapes.title
            shape.left, shape.top = Inches(MARGIN + 0.2), Inches(2.7)
            shape.width, shape.height = Inches(SLIDE_W - 2 * MARGIN - 0.4), Inches(1.6)
            shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            shape.text_frame.margin_left = 0
            para = shape.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT
            run = para.add_run()
            run.text = heading
            _style_run(run, 42, colors["title"], bold=True)
            if spec.get("subtitle"):
                _text(slide, Inches(MARGIN + 0.2), Inches(4.4), Inches(10), Inches(1), str(spec["subtitle"]), 22, colors["muted"])
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            _bg(slide, colors)
            _title(slide, heading or " ", colors)
            top, height = Inches(2.05), Inches(4.7)
            full_w = Inches(SLIDE_W - 2 * MARGIN)
            if kind == "bullets":
                items = list(spec.get("bullets") or [])
                if spec.get("body"):
                    if items:
                        items.insert(0, str(spec["body"]))
                    else:
                        _text(slide, Inches(MARGIN), top, full_w, height, str(spec["body"]), 24, colors["title"])
                if items:
                    _add_bullets(slide, Inches(MARGIN), top, full_w, height, items, colors)
            elif kind == "image":
                image = Path(str(spec.get("image") or "")).expanduser()
                items = list(spec.get("bullets") or [])
                if image.is_file():
                    if items:
                        _add_bullets(slide, Inches(MARGIN), top, Inches(5.2), height, items, colors, size=20)
                        _fit_picture(slide, str(image), Inches(6.3), top, Inches(6.2), height - Inches(0.4))
                    else:
                        _fit_picture(slide, str(image), Inches(MARGIN), top, full_w, height - Inches(0.4))
                    if spec.get("caption"):
                        _text(slide, Inches(MARGIN), Inches(6.55), full_w, Inches(0.4), str(spec["caption"]), 13,
                              colors["muted"], italic=True, align=PP_ALIGN.CENTER)
                else:
                    problems.append(f"slide {index + 1}: image not found ({spec.get('image')})")
                    _text(slide, Inches(MARGIN), top, full_w, Inches(1), f"[image not found: {spec.get('image')}]", 18, colors["muted"])
            elif kind == "table":
                rows = [[str(c) for c in row] for row in (spec.get("table") or []) if row]
                if rows:
                    _table(slide, rows, Inches(MARGIN), top, full_w, height, colors)
                else:
                    problems.append(f"slide {index + 1}: the table had no rows")
            elif kind == "quote":
                _text(slide, Inches(MARGIN + 0.4), Inches(2.3), Inches(SLIDE_W - 2 * MARGIN - 0.8), Inches(3.2),
                      "“" + str(spec.get("quote")).strip("“”\"") + "”", 32, colors["title"], italic=True,
                      anchor=MSO_ANCHOR.MIDDLE)
                if spec.get("attribution"):
                    _text(slide, Inches(MARGIN + 0.4), Inches(5.6), Inches(SLIDE_W - 2 * MARGIN - 0.8), Inches(0.5),
                          "— " + str(spec["attribution"]), 18, colors["muted"], align=PP_ALIGN.RIGHT)
            elif kind == "two_column":
                col_w = Inches((SLIDE_W - 2 * MARGIN - 0.5) / 2)
                for offset, key, head in ((0, "left", spec.get("left_title")), (1, "right", spec.get("right_title"))):
                    x = Inches(MARGIN) + offset * (col_w + Inches(0.5))
                    y = top
                    if head:
                        _text(slide, x, y, col_w, Inches(0.5), str(head), 22, colors["accent"], bold=True)
                        y += Inches(0.65)
                    items = list(spec.get(key) or [])
                    if items:
                        _add_bullets(slide, x, y, col_w, height - (y - top), items, colors, size=20)
            _footer(slide, index + 1, deck_title, colors)
        if kind in ("title", "section"):
            pass
        notes = str(spec.get("notes") or "").strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    file.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(file))
    except PermissionError:
        return {"status": "error", "retry_safe": False,
                "error": f"{file.name} is open in PowerPoint, so it can't be written. Close it and try again."}

    try:
        from tools.documents import pages
        back = pages.info(str(file))
    except Exception as exc:
        return {"status": "error", "retry_safe": False,
                "error": f"{file.name} was written but doesn't read back properly ({exc}). Don't hand it over."}
    result = {"status": "success", "path": str(file), "slides": back["count"], "theme": (theme or "graphite").lower(),
              "replaced": existed,
              "result": f"Made {file.name}: {back['count']} slides in the {(theme or 'graphite').lower()} look."}
    if problems:
        result["problems"] = problems
    return result

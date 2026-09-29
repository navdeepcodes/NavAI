"""Long documents, a page at a time: outline, reading, and search.

The old reader returned the first 12,000 characters of a file and stopped, so a
lecture PDF, a thesis or a textbook chapter was something Mike could only ever
read the start of. This treats every document as pages -- a PDF's pages, a
deck's slides, a Word file's sections, a text file's parts -- so Mike can see
its outline, read the pages that matter, and search across it (or a whole
folder of papers) for where something is said, quoting page numbers a student
can cite.

Pages with no text layer (scans, photos of notes) are read by OCR on the way,
transparently (tools/documents/ocr.py), and what OCR read is remembered on disk
so it's only paid for once.
"""
from __future__ import annotations

import hashlib
import json
import logging
import re
import time
from collections import OrderedDict
from pathlib import Path

from hostplatform import storage
from logs.logger import logger
from tools.documents import ocr
from tools.filesystem.document_reader import DocumentUnreadable
from tools.filesystem.path_utils import resolve_path

logging.getLogger("pypdf").setLevel(logging.ERROR)

#: How much of a document one read returns: a few thousand tokens.
READ_CHARS = 12_000
#: A page with fewer characters than this has no text layer worth the name.
MIN_TEXT = 20
#: Reading a scanned document by OCR stops after this long, and says where.
OCR_READ_BUDGET = 45.0
#: A whole-folder search stops after this long, and says how far it got.
SEARCH_BUDGET = 25.0
#: Files a folder search looks at, most recently changed first.
SEARCH_MAX_FILES = 300

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
_DOC_EXTS = {".pdf", ".docx", ".pptx"}
_CHUNK = 3000            # characters in a "part" of a text file or an unheaded Word file
_MAX_SECTION = 6000      # a Word section longer than this is split


def _kind(file: Path) -> str:
    ext = file.suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext == ".pptx":
        return "pptx"
    if ext in IMAGE_EXTS:
        return "image"
    return "text"


# ── one document ─────────────────────────────────────────────

class Doc:
    """A document as numbered pages. Built once, kept while it's in use."""

    def __init__(self, file: Path) -> None:
        self.file = file
        self.kind = _kind(file)
        self.title = self.author = ""
        self.outline: list[tuple[str, int]] = []      # (heading, page)
        self.labels: dict[int, str] = {}              # PDF page -> printed page label
        self.unit = {"pdf": "page", "pptx": "slide", "docx": "section",
                     "image": "image", "text": "part"}[self.kind]
        self._parts: list[tuple[str, str]] = []       # (label, text): docx, pptx, text
        self._reader = None
        self._texts: dict[int, str] = {}              # PDF layer text, as extracted
        self._ocr: dict[int, str] = {}                # pages read by OCR
        self._ocr_loaded = False
        self.count = 0
        loader = {"pdf": self._load_pdf, "docx": self._load_docx, "pptx": self._load_pptx,
                  "image": self._load_image, "text": self._load_text}[self.kind]
        loader()

    # -- loading --------------------------------------------------------
    def _load_pdf(self) -> None:
        from pypdf import PdfReader
        try:
            reader = PdfReader(str(self.file))
            if reader.is_encrypted and not reader.decrypt(""):
                raise DocumentUnreadable(
                    f"{self.file.name} is password-protected, so I can't read it. If the "
                    "student knows the password they can save an unlocked copy.")
            self.count = len(reader.pages)
        except DocumentUnreadable:
            raise
        except Exception as exc:
            raise DocumentUnreadable(f"{self.file.name} isn't a PDF I can open ({exc}).")
        self._reader = reader
        meta = reader.metadata
        if meta:
            self.title = str(getattr(meta, "title", "") or "").strip()
            self.author = str(getattr(meta, "author", "") or "").strip()
        try:
            for i, label in enumerate(reader.page_labels, start=1):
                if str(label) != str(i):
                    self.labels[i] = str(label)
        except Exception:
            pass
        try:
            self._walk_outline(reader.outline, reader)
        except Exception:
            logger.debug("Couldn't read the PDF outline.", exc_info=True)

    def _walk_outline(self, items, reader, depth: int = 0) -> None:
        for item in items:
            if isinstance(item, list):
                self._walk_outline(item, reader, depth + 1)
                continue
            try:
                number = reader.get_destination_page_number(item) + 1
            except Exception:
                continue
            title = str(getattr(item, "title", "") or "").strip()
            if title:
                self.outline.append(("  " * depth + title, number))

    def _load_docx(self) -> None:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        try:
            doc = Document(str(self.file))
        except Exception as exc:
            raise DocumentUnreadable(f"{self.file.name} isn't a Word file I can open ({exc}).")
        core = doc.core_properties
        self.title, self.author = (core.title or "").strip(), (core.author or "").strip()
        blocks: list[tuple[int, str]] = []               # (heading level or 0, text)
        for child in doc.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = Paragraph(child, doc)
                text = para.text.strip()
                if not text:
                    continue
                style = (para.style.name if para.style is not None else "") or ""
                level = 0
                match = re.match(r"Heading (\d)", style)
                if match:
                    level = int(match.group(1))
                elif style == "Title":
                    level = 1
                blocks.append((level, text))
            elif tag == "tbl":
                rows = [" | ".join(c.text.strip().replace("\n", " ") for c in row.cells)
                        for row in Table(child, doc).rows]
                blocks.append((0, "\n".join(rows)))
        self._parts = _sections(blocks)
        self.count = len(self._parts)
        self.outline = [(label, i) for i, (label, _t) in enumerate(self._parts, start=1)
                        if not label.startswith("Part ")]

    def _load_pptx(self) -> None:
        from pptx import Presentation
        try:
            deck = Presentation(str(self.file))
        except Exception as exc:
            raise DocumentUnreadable(f"{self.file.name} isn't a PowerPoint file I can open ({exc}).")
        core = deck.core_properties
        self.title, self.author = (core.title or "").strip(), (core.author or "").strip()
        for i, slide in enumerate(deck.slides, start=1):
            title = ""
            title_shape = None
            try:
                title_shape = slide.shapes.title
                if title_shape is not None and title_shape.has_text_frame:
                    title = title_shape.text_frame.text.strip()
            except Exception:
                pass
            lines: list[str] = []
            for shape in slide.shapes:
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if shape.has_text_frame:
                    lines.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                elif getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        lines.append(" | ".join(c.text.strip() for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"[Speaker notes] {notes}")
            label = f"Slide {i}: {title}" if title else f"Slide {i}"
            self._parts.append((label, "\n".join(([title] if title else []) + lines)))
            if title:
                self.outline.append((title, i))
        self.count = len(self._parts)

    def _load_image(self) -> None:
        self.count = 1
        self._parts = [("Image", "")]

    def _load_text(self) -> None:
        raw = _read_text_file(self.file)
        if raw is None:
            raise DocumentUnreadable(
                f"{self.file.name} isn't a text document Mike can read. Mike reads PDF, Word "
                "(.docx), PowerPoint (.pptx), plain text and pictures of text.")
        lines = raw.splitlines()
        part, size, first = [], 0, 1
        for number, line in enumerate(lines, start=1):
            part.append(line)
            size += len(line) + 1
            if size >= _CHUNK:
                self._parts.append((f"lines {first}-{number}", "\n".join(part)))
                part, size, first = [], 0, number + 1
        if part or not self._parts:
            self._parts.append((f"lines {first}-{max(first, len(lines))}", "\n".join(part)))
        self.count = len(self._parts)

    # -- text -----------------------------------------------------------
    def layer_text(self, n: int) -> str:
        """The text the file itself carries for page n."""
        if self.kind == "pdf":
            if n not in self._texts:
                try:
                    self._texts[n] = (self._reader.pages[n - 1].extract_text() or "").strip()
                except Exception:
                    logger.debug("Couldn't extract PDF page %s.", n, exc_info=True)
                    self._texts[n] = ""
            return self._texts[n]
        return self._parts[n - 1][1] if 1 <= n <= len(self._parts) else ""

    def _ocr_file(self) -> Path:
        st = self.file.stat()
        key = hashlib.sha1(f"{self.file}|{st.st_mtime_ns}|{st.st_size}".encode()).hexdigest()[:24]
        return storage.data_dir() / "doc_cache" / f"{key}.json"

    def _load_ocr(self) -> None:
        if self._ocr_loaded:
            return
        self._ocr_loaded = True
        try:
            data = json.loads(self._ocr_file().read_text(encoding="utf-8"))
            self._ocr = {int(k): str(v) for k, v in (data.get("pages") or {}).items()}
        except (OSError, ValueError):
            pass

    def _save_ocr(self) -> None:
        try:
            target = self._ocr_file()
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(json.dumps({"file": self.file.name,
                                          "pages": {str(k): v for k, v in self._ocr.items()}}),
                              encoding="utf-8")
        except OSError:
            logger.debug("Couldn't save the OCR cache.", exc_info=True)

    def needs_ocr(self, n: int) -> bool:
        """Does this page have only a picture of its text, not yet read?"""
        if self.kind not in ("pdf", "image"):
            return False
        self._load_ocr()
        if n in self._ocr:
            return False
        return self.kind == "image" or len(self.layer_text(n)) < MIN_TEXT

    def read_by_ocr(self, pages: list[int], deadline: float) -> list[int]:
        """OCR these pages (in batches, stopping at the deadline); the pages
        that were done."""
        done: list[int] = []
        todo = [n for n in pages if self.needs_ocr(n)]
        for start in range(0, len(todo), 6):
            if time.monotonic() > deadline:
                break
            batch = todo[start:start + 6]
            if self.kind == "image":
                texts = {1: ocr.tidy(ocr.read_images([str(self.file)])[0])}
            else:
                texts = ocr.read_pdf_pages(str(self.file), batch)
            self._ocr.update(texts)
            done.extend(texts)
        if done:
            self._save_ocr()
        return done

    def text(self, n: int) -> tuple[str, str]:
        """(text, how): how is "text" (the file's own), "ocr", or "none"."""
        layer = self.layer_text(n)
        if self.kind not in ("pdf", "image") or len(layer) >= MIN_TEXT:
            return layer, "text"
        self._load_ocr()
        read = self._ocr.get(n, "")
        if read:
            return read, "ocr"
        return layer, "none"

    def label(self, n: int) -> str:
        if self.kind == "pdf":
            extra = f" (printed {self.labels[n]})" if n in self.labels else ""
            return f"Page {n}{extra}"
        if self.kind == "pptx":
            return self._parts[n - 1][0]
        if self.kind == "docx":
            return f"Section {n}: {self._parts[n - 1][0]}"
        if self.kind == "image":
            return "Image"
        return f"Part {n} ({self._parts[n - 1][0]})"


def _read_text_file(file: Path) -> str | None:
    try:
        data = file.read_bytes()
    except OSError:
        return None
    if b"\x00" in data[:4096]:
        return None
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return None


def _sections(blocks: list[tuple[int, str]]) -> list[tuple[str, str]]:
    """A Word file's blocks as sections: split at the headings that give three
    or more, else in even parts. (label, text) each."""
    levels = sorted({lv for lv, _ in blocks if lv})
    split_at = 0
    for level in levels:
        if sum(1 for lv, _ in blocks if 0 < lv <= level) >= 3:
            split_at = level
            break
    sections: list[tuple[str, list[str]]] = []
    if split_at:
        title, lines = "Start", []
        for level, text in blocks:
            if 0 < level <= split_at:
                if lines:
                    sections.append((title, lines))
                title, lines = text[:80], [text]
            else:
                lines.append(("#" * level + " " if level else "") + text)
        if lines:
            sections.append((title, lines))
    else:
        title, lines, size = "Part 1", [], 0
        for _level, text in blocks:
            lines.append(text)
            size += len(text) + 1
            if size >= _CHUNK:
                sections.append((title, lines))
                title, lines, size = f"Part {len(sections) + 1}", [], 0
        if lines or not sections:
            sections.append((title, lines))
    out: list[tuple[str, str]] = []
    for title, lines in sections:
        text = "\n".join(lines)
        if len(text) <= _MAX_SECTION:
            out.append((title, text))
            continue
        pieces, buffer, size = [], [], 0
        for line in lines:
            buffer.append(line)
            size += len(line) + 1
            if size >= _CHUNK:
                pieces.append("\n".join(buffer))
                buffer, size = [], 0
        if buffer:
            pieces.append("\n".join(buffer))
        for i, piece in enumerate(pieces, start=1):
            out.append((title if i == 1 else f"{title} (continued)", piece))
    return out


# ── opening documents ────────────────────────────────────────

_docs: "OrderedDict[tuple, Doc]" = OrderedDict()
MAX_OPEN = 6


def open_doc(path: str) -> Doc:
    file = resolve_path(path)
    if not file.exists():
        raise FileNotFoundError(f"File not found: {file}")
    if file.is_dir():
        raise DocumentUnreadable(f"{file} is a folder. Give a file, or use search_document to look through a folder.")
    stat = file.stat()
    key = (str(file), stat.st_mtime_ns, stat.st_size)
    if key in _docs:
        _docs.move_to_end(key)
        return _docs[key]
    doc = Doc(file)
    _docs[key] = doc
    while len(_docs) > MAX_OPEN:
        _docs.popitem(last=False)
    return doc


def parse_pages(spec, count: int) -> list[int]:
    """"3", "3-7", "1,4,9-11", "12-" (to the end), "last", "all" -> page numbers.
    Raises ValueError, saying what's wrong, for anything out of range."""
    if spec is None or (isinstance(spec, str) and not spec.strip()):
        return []
    if isinstance(spec, (int, float)):
        spec = str(int(spec))
    if isinstance(spec, (list, tuple)):
        spec = ",".join(str(s) for s in spec)
    text = str(spec).strip().lower()
    if text == "all":
        return list(range(1, count + 1))
    text = re.sub(r"\s*(?:\bto\b|–|—)\s*", "-", text)         # "2 to 4", "2 – 4"
    pages: list[int] = []
    for part in re.split(r"[,\s]+", text):
        if not part:
            continue
        if part == "last":
            pages.append(count)
            continue
        m = re.fullmatch(r"(\d+)\s*(?:-|–|to)\s*(\d*)", part)
        if m:
            start = int(m.group(1))
            end = int(m.group(2)) if m.group(2) else count
            if start > end:
                raise ValueError(f"'{part}' runs backwards.")
            pages.extend(range(start, end + 1))
        elif part.isdigit():
            pages.append(int(part))
        else:
            raise ValueError(f"I can't read '{part}' as a page. Use forms like 3, 3-7 or 1,4,9-11.")
    bad = [p for p in pages if not 1 <= p <= count]
    if bad:
        raise ValueError(f"Page {bad[0]} isn't there: it has {count} {'page' if count == 1 else 'pages'}.")
    seen: set[int] = set()
    return [p for p in pages if not (p in seen or seen.add(p))]


def ranges(pages: list[int]) -> str:
    """[1, 2, 3, 7, 9, 10] -> "1-3,7,9-10"."""
    if not pages:
        return ""
    out, start, prev = [], pages[0], pages[0]
    for p in pages[1:]:
        if p == prev + 1:
            prev = p
            continue
        out.append(f"{start}-{prev}" if prev > start else str(start))
        start = prev = p
    out.append(f"{start}-{prev}" if prev > start else str(start))
    return ",".join(out)


# ── the three things Mike does with a document ───────────────

def info(path: str) -> dict:
    """What a document is and how it's laid out: pages, title, outline, and
    whether it's scanned -- so Mike knows where to read before reading."""
    doc = open_doc(path)
    result: dict = {"status": "success", "path": str(doc.file), "kind": doc.kind,
                    "unit": doc.unit, "count": doc.count,
                    "size_kb": round(doc.file.stat().st_size / 1024)}
    if doc.title:
        result["title"] = doc.title
    if doc.author:
        result["author"] = doc.author
    if doc.outline:
        shown = doc.outline[:60]
        result["outline"] = [("  " if title.startswith(" ") else "") + f"{title.strip()} — {doc.unit} {page}"
                             for title, page in shown]
        if len(doc.outline) > len(shown):
            result["outline_more"] = len(doc.outline) - len(shown)
    if doc.kind in ("pdf", "image"):
        sample = list(range(1, doc.count + 1))
        if len(sample) > 12:
            step = len(sample) / 12
            sample = sorted({int(i * step) + 1 for i in range(12)})
        empty = [n for n in sample if doc.needs_ocr(n) or doc.text(n)[1] == "none"]
        result["scanned"] = bool(sample) and len(empty) >= max(1, len(sample) // 2)
        if empty:
            result["pages_without_text_in_sample"] = ranges(empty)
        if result["scanned"]:
            result["note"] = ("Mostly pictures of text. " + (
                "Reading it uses OCR (about a second a page), so read the pages you need."
                if ocr.available() else f"OCR isn't available here: {ocr.unavailable_reason()}."))
    result["result"] = (f"{doc.file.name}: {doc.count} {doc.unit}{'s' if doc.count != 1 else ''}"
                        + (f", \"{doc.title}\"" if doc.title else "") + ".")
    return result


def read(path: str, pages=None, max_chars: int = READ_CHARS) -> dict:
    """The text of some pages: those asked for, or from the start until about
    max_chars characters. Says what came next, and what needed OCR."""
    doc = open_doc(path)
    try:
        wanted = parse_pages(pages, doc.count)
    except ValueError as exc:
        return {"status": "error", "error": str(exc), "retry_safe": True, "total": doc.count}
    explicit = bool(wanted)
    order = wanted or list(range(1, doc.count + 1))
    deadline = time.monotonic() + OCR_READ_BUDGET
    shown: list[int] = []
    ocr_pages: list[int] = []
    unread: list[int] = []
    chunks: list[str] = []
    used = 0
    cut_page = None

    index = 0
    while index < len(order):
        n = order[index]
        if doc.needs_ocr(n) and ocr.available():
            if time.monotonic() > deadline:
                unread = order[index:]
                break
            doc.read_by_ocr([p for p in order[index:index + 6] if doc.needs_ocr(p)], deadline)
        text, how = doc.text(n)
        header = f"--- {doc.label(n)} ---"
        body = text if text.strip() else (
            "(no text could be read from this page)" if how == "none" else "(no text on this page)")
        size = len(header) + len(body) + 2
        if used and used + size > max_chars:
            break
        if size > max_chars:
            body = body[: max_chars - len(header) - 60].rstrip() + "\n… (this page is longer; the rest was cut)"
            cut_page = n
        chunks.append(f"{header}\n{body}")
        used += size
        shown.append(n)
        if how == "ocr":
            ocr_pages.append(n)
        index += 1
        if used >= max_chars:
            break

    result: dict = {"status": "success", "path": str(doc.file), "kind": doc.kind, "unit": doc.unit,
                    "total": doc.count, "shown": ranges(shown), "text": "\n\n".join(chunks)}
    remaining = [p for p in order if p not in shown]
    if remaining and not unread:
        result["next"] = ranges(remaining if explicit else remaining[: max(1, len(shown))])
    notes: list[str] = []
    if ocr_pages:
        notes.append(f"{doc.unit.capitalize()} {ranges(ocr_pages)} had no text of their own, so I read them "
                     "by OCR: it can misread words, numbers and anything handwritten.")
    if unread:
        notes.append(f"Ran out of time reading pictures of text: {doc.unit}s {ranges(unread)} weren't read yet. "
                     "Ask for them again to carry on.")
    blank = [n for n in shown if doc.text(n)[1] == "none"]
    if blank and not ocr.available() and doc.kind in ("pdf", "image"):
        notes.append(f"{doc.unit.capitalize()} {ranges(blank)} are pictures of text and OCR isn't available "
                     f"here ({ocr.unavailable_reason()}).")
    if cut_page:
        notes.append(f"{doc.unit.capitalize()} {cut_page} is very long and was cut; "
                     "search_document can find the part you need.")
    if notes:
        result["notes"] = notes
    if shown and not any(doc.text(n)[0].strip() for n in shown):
        result["status"] = "error"
        result["retry_safe"] = False
        result["error"] = f"No text could be read from {doc.file.name}" + (
            f": it looks like a scan and {ocr.unavailable_reason()}."
            if doc.kind in ("pdf", "image") and not ocr.available() else ".")
    return result


_STOP = frozenset("the a an and or of to in on for with is are was were be by at as it this that from".split())


def _snippets(text: str, matches: list, width: int = 90, limit: int = 2) -> list[str]:
    out = []
    for m in matches[:limit]:
        start, end = max(0, m.start() - width), min(len(text), m.end() + width)
        piece = re.sub(r"\s+", " ", text[start:end]).strip()
        out.append(("…" if start else "") + piece + ("…" if end < len(text) else ""))
    return out


def _files_under(root: Path) -> list[Path]:
    found = []
    for item in root.rglob("*"):
        try:
            if any(part.startswith(".") for part in item.relative_to(root).parts):
                continue
            ext = item.suffix.lower()
            if item.is_file() and (ext in _DOC_EXTS or ext in IMAGE_EXTS or ext in (".txt", ".md", ".tex", ".rst")):
                found.append(item)
        except OSError:
            continue
    found.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    return found


def search(query: str, path: str = ".", regex: bool = False, max_results: int = 15,
           budget: float = SEARCH_BUDGET) -> dict:
    """Where a document -- or every document in a folder -- says something.
    Page by page: file, page, the matching lines. A phrase first; if nothing
    has the phrase, the pages that have all of its words."""
    query = (query or "").strip()
    if not query:
        return {"status": "error", "error": "Say what to look for.", "retry_safe": True}
    root = resolve_path(path or ".")
    if not root.exists():
        return {"status": "error", "error": f"Nothing at {root}.", "retry_safe": True}
    files = [root] if root.is_file() else _files_under(root)[:SEARCH_MAX_FILES]
    flags = 0 if any(c.isupper() for c in query) and not regex else re.IGNORECASE
    try:
        # a phrase matches across line breaks and the odd spacing extraction
        # leaves (PDF text breaks at every bold or italic change)
        phrase = re.compile(query if regex else r"\s+".join(re.escape(w) for w in query.split()), flags)
    except re.error as exc:
        return {"status": "error", "error": f"That isn't a valid pattern: {exc}", "retry_safe": True}
    words = [] if regex else [w for w in re.findall(r"\w+", query.lower()) if len(w) > 2 and w not in _STOP]
    word_res = [re.compile(r"\b" + re.escape(w) + r"\b", re.IGNORECASE) for w in dict.fromkeys(words)]

    started = time.monotonic()
    deadline = started + budget
    phrase_hits: list[dict] = []
    word_hits: list[dict] = []
    searched_pages = searched_files = unsearched_scans = 0
    skipped: list[str] = []
    stopped_early = False

    for file in files:
        if time.monotonic() > deadline:
            stopped_early = True
            break
        try:
            doc = open_doc(str(file))
        except (DocumentUnreadable, FileNotFoundError, OSError) as exc:
            skipped.append(f"{file.name}: {str(exc)[:80]}")
            continue
        searched_files += 1
        name = file.name if root.is_file() else str(file.relative_to(root))
        pages = list(range(1, doc.count + 1))
        if doc.kind in ("pdf", "image") and ocr.available():
            doc.read_by_ocr([n for n in pages if doc.needs_ocr(n)], deadline)
        for n in pages:
            if time.monotonic() > deadline:
                stopped_early = True
                break
            text, how = doc.text(n)
            searched_pages += 1
            if how == "none" and doc.kind in ("pdf", "image"):
                unsearched_scans += 1
                continue
            found = list(phrase.finditer(text))
            entry = None
            if found:
                entry = {"file": name, "page": n, "label": doc.label(n), "matches": len(found),
                         "snippets": _snippets(text, found)}
                phrase_hits.append(entry)
            elif len(word_res) >= 2 and all(w.search(text) for w in word_res):
                first = word_res[0].search(text)
                entry = {"file": name, "page": n, "label": doc.label(n),
                         "matches": sum(len(w.findall(text)) for w in word_res),
                         "snippets": _snippets(text, [first]) if first else []}
                word_hits.append(entry)
            if entry and how == "ocr":
                entry["by_ocr"] = True
        if stopped_early:
            break

    hits = phrase_hits or word_hits
    hits.sort(key=lambda h: -h["matches"])
    shown = hits[: max(1, int(max_results or 15))]
    took = time.monotonic() - started
    summary = (f"Searched {searched_files} file(s), {searched_pages} page(s) in {took:.0f}s: "
               f"{len(hits)} page(s) match"
               + (" -- none has the exact phrase, so these have all its words" if word_hits and not phrase_hits else "")
               + ".")
    result: dict = {"status": "success", "query": query,
                    "matched_on": ("phrase" if phrase_hits else "words") if hits else "nothing",
                    "result": summary, "hits": shown}
    if len(hits) > len(shown):
        result["more"] = len(hits) - len(shown)
    notes = []
    if stopped_early:
        notes.append(f"Stopped after {budget:.0f}s with more still to search: narrow it to a folder or ask again.")
    if unsearched_scans:
        notes.append(f"{unsearched_scans} page(s) are pictures of text that couldn't be read"
                     + ("" if ocr.available() else f" ({ocr.unavailable_reason()})") + ", so they weren't searched.")
    if skipped:
        notes.append("Couldn't open: " + "; ".join(skipped[:4]))
    if notes:
        result["notes"] = notes
    return result

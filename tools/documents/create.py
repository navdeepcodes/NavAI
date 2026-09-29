"""Making documents: an essay, a report, a resume, a lab write-up -- as Word or PDF.

Mike writes the content as Markdown (headings, lists, tables, code, bold and
italic, images, page breaks) and this makes the file. The same content gives a
.docx (opens and edits in Word) or a PDF (ready to hand in), formatted by a
preset -- MLA, APA, report, resume, letter, or plain -- that is only settings:
font, size, line spacing, margins, page numbers, indents. Anything can be
overridden per call, because a course's own rules win.

Written files are read back and summarised from what is actually in them, not
from what was asked for.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field, replace
from pathlib import Path

from logs.logger import logger

FENCE = "```"


# ── style presets: settings, not rules ───────────────────────

@dataclass(frozen=True)
class Style:
    font: str = "Calibri"
    size: float = 11.0
    line_spacing: float = 1.15            # 1.0 single, 2.0 double
    space_after: float = 6.0              # points after a paragraph
    margin_in: float = 1.0
    page_size: str = "letter"             # letter | a4
    page_numbers: str = "bottom-center"   # bottom-center | top-right | none
    header_name: bool = False             # "Surname 3": the running head MLA asks for
    first_line_indent_in: float = 0.0
    title_page: bool = False
    center_title: bool = False
    heading_font: str = ""
    heading_color: str = "1F1F23"
    heading_sizes: tuple = (20.0, 15.0, 12.5)
    heading_bold: bool = True
    heading_center_h1: bool = False
    rule_under_h2: bool = False
    hanging_bibliography: bool = False


PRESETS: dict[str, Style] = {
    "plain": Style(),
    "report": Style(font="Calibri", size=11, line_spacing=1.2, title_page=True),
    "mla": Style(font="Times New Roman", size=12, line_spacing=2.0, space_after=0,
                 first_line_indent_in=0.5, page_numbers="top-right", header_name=True,
                 center_title=True, heading_font="Times New Roman", heading_color="000000",
                 heading_sizes=(12, 12, 12), heading_bold=False, heading_center_h1=True,
                 hanging_bibliography=True),
    "apa": Style(font="Times New Roman", size=12, line_spacing=2.0, space_after=0,
                 first_line_indent_in=0.5, page_numbers="top-right", title_page=True,
                 center_title=True, heading_font="Times New Roman", heading_color="000000",
                 heading_sizes=(12, 12, 12), heading_bold=True, heading_center_h1=True,
                 hanging_bibliography=True),
    "resume": Style(font="Calibri", size=10.5, line_spacing=1.0, space_after=2, margin_in=0.6,
                    page_numbers="none", center_title=True, heading_sizes=(22, 11.5, 10.5),
                    rule_under_h2=True),
    "letter": Style(font="Calibri", size=11, line_spacing=1.0, space_after=10, page_numbers="none",
                    heading_sizes=(14, 12, 11)),
}

_PAGE_IN = {"letter": (8.5, 11.0), "a4": (8.27, 11.69)}


def style_for(name: str | None, **override) -> Style:
    base = PRESETS.get((name or "plain").lower(), PRESETS["plain"])
    clean = {k: v for k, v in override.items() if v not in (None, "")}
    return replace(base, **clean) if clean else base


# ── Markdown -> blocks ───────────────────────────────────────

@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    link: str = ""


@dataclass
class Block:
    kind: str                       # h p ul ol table code quote hr image pagebreak
    level: int = 0
    runs: list = field(default_factory=list)        # inline runs (h, p, quote)
    items: list = field(default_factory=list)       # (depth, runs) for lists
    rows: list = field(default_factory=list)        # table: list of list of runs; row 0 is the header
    text: str = ""                                  # code text / image path
    alt: str = ""
    aligns: list = field(default_factory=list)


_INLINE = re.compile(
    r"(\*\*\*(?P<bi>.+?)\*\*\*)|(\*\*(?P<b>.+?)\*\*)|(__(?P<b2>.+?)__)|(\*(?P<i>[^*\s][^*]*?)\*)"
    r"|(?<![\w])_(?P<i2>[^_\s][^_]*?)_(?![\w])|(`(?P<c>[^`]+)`)|(\[(?P<lt>[^\]]+)\]\((?P<lu>[^)\s]+)\))")


def inline(text: str) -> list[Run]:
    runs: list[Run] = []
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            runs.append(Run(text[pos:m.start()]))
        if m.group("bi") is not None:
            runs.append(Run(m.group("bi"), bold=True, italic=True))
        elif m.group("b") is not None or m.group("b2") is not None:
            runs.append(Run(m.group("b") or m.group("b2"), bold=True))
        elif m.group("i") is not None or m.group("i2") is not None:
            runs.append(Run(m.group("i") or m.group("i2"), italic=True))
        elif m.group("c") is not None:
            runs.append(Run(m.group("c"), code=True))
        else:
            runs.append(Run(m.group("lt"), link=m.group("lu")))
        pos = m.end()
    if pos < len(text):
        runs.append(Run(text[pos:]))
    return [r for r in runs if r.text] or [Run("")]


_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def clean(text: str) -> str:
    """Characters a Word file can't hold (form feeds and the like, which text
    taken from a PDF is full of) removed; a form feed was a page break."""
    return _XML_BAD.sub("", (text or "").replace("\x0c", "\n\n"))


def parse(markdown: str) -> list[Block]:
    lines = clean(markdown).replace("\r\n", "\n").replace("\t", "    ").split("\n")
    blocks: list[Block] = []
    i = 0
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append(Block("p", runs=inline(" ".join(s.strip() for s in para))))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith(FENCE):
            flush()
            lang, body = stripped[3:].strip(), []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(FENCE):
                body.append(lines[i])
                i += 1
            blocks.append(Block("code", text="\n".join(body), alt=lang))
            i += 1
            continue
        if not stripped:
            flush()
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", stripped)
        if heading:
            flush()
            blocks.append(Block("h", level=min(3, len(heading.group(1))), runs=inline(heading.group(2))))
            i += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            flush()
            blocks.append(Block("hr"))
            i += 1
            continue
        if stripped.lower() in ("\\pagebreak", "<pagebreak>", "[pagebreak]", "\\newpage", "---pagebreak---"):
            flush()
            blocks.append(Block("pagebreak"))
            i += 1
            continue
        image = re.fullmatch(r"!\[(.*?)\]\((.+?)\)", stripped)
        if image:
            flush()
            blocks.append(Block("image", alt=image.group(1), text=image.group(2)))
            i += 1
            continue
        if "|" in stripped and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]) and "-" in lines[i + 1]:
            flush()
            header = _cells(stripped)
            aligns = []
            for cell in _cells(lines[i + 1]):
                aligns.append("center" if cell.startswith(":") and cell.endswith(":")
                              else "right" if cell.endswith(":") else "left")
            rows = [[inline(c) for c in header]]
            i += 2
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                cells = (_cells(lines[i]) + [""] * len(header))[:len(header)]
                rows.append([inline(c) for c in cells])
                i += 1
            blocks.append(Block("table", rows=rows, aligns=aligns))
            continue
        if stripped.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(Block("quote", runs=inline(" ".join(quote))))
            continue
        bullet = re.match(r"^(\s*)([-*+•])\s+(.*)$", line)
        number = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if bullet or number:
            flush()
            ordered = bool(number)
            pattern = r"^(\s*)(\d+)[.)]\s+(.*)$" if ordered else r"^(\s*)([-*+•])\s+(.*)$"
            items = []
            while i < len(lines):
                m = re.match(pattern, lines[i])
                if not m:
                    # a wrapped continuation of the previous item
                    if lines[i].strip() and lines[i].startswith("  ") and items:
                        depth, runs = items[-1]
                        items[-1] = (depth, runs + inline(" " + lines[i].strip()))
                        i += 1
                        continue
                    break
                items.append((len(m.group(1)) // 2, inline(m.group(3))))
                i += 1
            blocks.append(Block("ol" if ordered else "ul", items=items))
            continue
        para.append(line)
        i += 1
    flush()
    return blocks


def plain(runs: list[Run]) -> str:
    return "".join(r.text for r in runs)


# ── Word ─────────────────────────────────────────────────────

_HEADING_WORDS = re.compile(r"(works cited|references|bibliography|sources)")
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _rgb(hexstr: str):
    from docx.shared import RGBColor
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(font, name: str, size: float | None = None, bold=None, italic=None, color: str = "") -> None:
    """Font name (for every script, so Word doesn't substitute), size, weight, colour."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    font.name = name
    rpr = font._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color:
        font.color.rgb = _rgb(color)


def _field(run, instruction: str, shown: str = "1") -> None:
    """A Word field (PAGE, TOC ...) that Word fills in and updates."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = shown
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, text, end):
        run._r.append(el)


def _shade(paragraph, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def _border(paragraph, side: str = "bottom", size: int = 6, color: str = "999999") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "1")
    edge.set(qn("w:color"), color)
    borders.append(edge)
    paragraph._p.get_or_add_pPr().append(borders)


def _runs(paragraph, runs: list[Run], font: str, size: float, bold=None, italic=None, color: str = "") -> None:
    for r in runs:
        run = paragraph.add_run(r.text)
        _set_font(run.font, "Consolas" if r.code else font, size - (1 if r.code else 0),
                  bold=True if (r.bold or bold) else None, italic=True if (r.italic or italic) else None,
                  color=color)
        if r.link:
            run.font.underline = True
            run.font.color.rgb = _rgb("1155CC")


def _spacing(paragraph, style: Style, after: float | None = None, indent: float | None = None,
             line: float | None = None) -> None:
    from docx.shared import Inches, Pt
    fmt = paragraph.paragraph_format
    fmt.line_spacing = style.line_spacing if line is None else line
    fmt.space_after = Pt(style.space_after if after is None else after)
    fmt.space_before = Pt(0)
    if indent:
        fmt.first_line_indent = Inches(indent)


def write_docx(path: Path, markdown: str, style: Style, title: str = "", author: str = "",
               toc: bool = False) -> dict:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt

    doc = Document()
    width, height = _PAGE_IN.get(style.page_size, _PAGE_IN["letter"])
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(width), Inches(height)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(style.margin_in))
    section.header_distance = Inches(0.5)

    _set_font(doc.styles["Normal"].font, style.font, style.size)
    heading_font = style.heading_font or style.font
    double = style.line_spacing >= 2
    for level in (1, 2, 3):
        hs = doc.styles[f"Heading {level}"]
        italic = level == 3 and style.heading_bold and style.heading_center_h1 and style.font == "Times New Roman"
        _set_font(hs.font, heading_font, style.heading_sizes[level - 1], bold=style.heading_bold,
                  italic=italic, color=style.heading_color)
        hs.paragraph_format.space_before = Pt(0 if double else (14 if level == 1 else 10))
        hs.paragraph_format.space_after = Pt(0 if double else 6)
        hs.paragraph_format.line_spacing = style.line_spacing if double else 1.1
        hs.paragraph_format.keep_with_next = True
    if title:
        doc.core_properties.title = title
    if author:
        doc.core_properties.author = author
    doc.core_properties.comments = "Made by Mike"

    if style.page_numbers != "none" or style.header_name:
        top = style.page_numbers == "top-right" or style.header_name
        para = (section.header if top else section.footer).paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if top else WD_ALIGN_PARAGRAPH.CENTER
        if style.header_name and author.strip():
            _set_font(para.add_run(author.strip().split()[-1] + " ").font, style.font, style.size)
        if style.page_numbers != "none":
            run = para.add_run()
            _set_font(run.font, style.font, style.size)
            _field(run, "PAGE")

    counts = {"words": 0, "h": 0, "table": 0, "image": 0, "code": 0, "list": 0}
    in_refs = False
    seen_h1 = False

    if style.title_page and title:
        for _ in range(7):
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, style, after=12, line=1.2)
        _set_font(p.add_run(title).font, heading_font, style.heading_sizes[0] + 4, bold=True,
                  color=style.heading_color)
        if author:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _spacing(p, style, after=4, line=1.2)
            _set_font(p.add_run(author).font, style.font, style.size)
        doc.add_page_break()
        seen_h1 = True
    if toc:
        p = doc.add_paragraph("Contents")
        p.style = doc.styles["Heading 1"]
        _field(doc.add_paragraph().add_run(), 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to fill this in.")
        doc.add_page_break()
        flag = OxmlElement("w:updateFields")
        flag.set(_W + "val", "true")
        doc.settings.element.append(flag)

    for b in parse(markdown):
        if b.kind == "h":
            counts["h"] += 1
            text = plain(b.runs)
            in_refs = bool(_HEADING_WORDS.fullmatch(text.strip().lower()))
            p = doc.add_paragraph()
            p.style = doc.styles[f"Heading {b.level}"]
            _runs(p, b.runs, heading_font, style.heading_sizes[b.level - 1], bold=style.heading_bold or None,
                  color=style.heading_color)
            if b.level == 1 and (style.heading_center_h1 or (style.center_title and not seen_h1)):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if style.rule_under_h2 and b.level == 2:
                _border(p, color="B9BAC0")
            seen_h1 = seen_h1 or b.level == 1
            counts["words"] += len(text.split())
        elif b.kind == "p":
            p = doc.add_paragraph()
            hang = in_refs and style.hanging_bibliography
            _spacing(p, style, indent=None if hang else (style.first_line_indent_in or None))
            if hang:
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
            _runs(p, b.runs, style.font, style.size)
            counts["words"] += len(plain(b.runs).split())
        elif b.kind in ("ul", "ol"):
            counts["list"] += 1
            for n, (depth, runs) in enumerate(b.items, start=1):
                p = doc.add_paragraph()
                fmt = p.paragraph_format
                left = 0.25 + 0.25 * (depth + 1)
                fmt.left_indent = Inches(left)
                fmt.first_line_indent = Inches(-0.25)
                fmt.line_spacing = style.line_spacing
                fmt.space_after = Pt(0 if double else min(style.space_after, 3))
                fmt.tab_stops.add_tab_stop(Inches(left))
                _set_font(p.add_run(("•" if b.kind == "ul" else f"{n}.") + "\t").font, style.font, style.size)
                _runs(p, runs, style.font, style.size)
                counts["words"] += len(plain(runs).split())
        elif b.kind == "table":
            counts["table"] += 1
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(b.rows):
                for c, cell_runs in enumerate(row):
                    cell = table.cell(r, c)
                    para = cell.paragraphs[0]
                    para.paragraph_format.space_after = Pt(2)
                    para.paragraph_format.line_spacing = 1.0
                    align = b.aligns[c] if c < len(b.aligns) else "left"
                    para.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER,
                                      "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(align, WD_ALIGN_PARAGRAPH.LEFT)
                    _runs(para, cell_runs, style.font, max(9, style.size - 1), bold=(r == 0) or None)
                    if r == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(_W + "val", "clear")
                        shade.set(_W + "fill", "E9E9EC")
                        cell._tc.get_or_add_tcPr().append(shade)
                    counts["words"] += len(plain(cell_runs).split())
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif b.kind == "code":
            counts["code"] += 1
            for line in b.text.split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Inches(0.15)
                _set_font(p.add_run(line or " ").font, "Consolas", max(8.5, style.size - 2))
                _shade(p, "F2F2F4")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif b.kind == "quote":
            p = doc.add_paragraph()
            _spacing(p, style)
            p.paragraph_format.left_indent = Inches(0.5)
            _runs(p, b.runs, style.font, style.size, italic=True)
            _border(p, side="left", size=12, color="B9BAC0")
            counts["words"] += len(plain(b.runs).split())
        elif b.kind == "hr":
            _border(doc.add_paragraph(), color="BBBBBB")
        elif b.kind == "pagebreak":
            doc.add_page_break()
        elif b.kind == "image":
            source = Path(b.text).expanduser()
            if not source.is_file():
                doc.add_paragraph(f"[image not found: {b.text}]")
                continue
            counts["image"] += 1
            try:
                from PIL import Image
                with Image.open(source) as im:
                    w_px = im.size[0]
                usable = width - 2 * style.margin_in
                doc.add_picture(str(source), width=Inches(min(usable, w_px / 96.0 if w_px else usable)))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if b.alt:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _spacing(cap, style, after=8, line=1.0)
                    _runs(cap, [Run(b.alt, italic=True)], style.font, style.size - 1)
            except Exception:
                logger.exception("Couldn't place the image %s", source)
                doc.add_paragraph(f"[image couldn't be placed: {b.text}]")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return counts


# ── PDF (Qt's own text layout and PDF writer) ────────────────

def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _html_runs(runs: list[Run]) -> str:
    out = []
    for r in runs:
        t = _esc(r.text)
        if r.code:
            t = f'<span style="font-family:Consolas,monospace; background-color:#f0f0f2;">{t}</span>'
        if r.bold:
            t = f"<b>{t}</b>"
        if r.italic:
            t = f"<i>{t}</i>"
        if r.link:
            t = f'<a href="{_esc(r.link)}" style="color:#1155cc;">{t}</a>'
        out.append(t)
    return "".join(out)


def _to_html(blocks: list[Block], style: Style, title: str, author: str) -> tuple[str, dict]:
    line = int(style.line_spacing * 100 + (5 if style.line_spacing < 1.5 else 0))
    counts = {"words": 0, "h": 0, "table": 0, "image": 0, "code": 0, "list": 0}
    heading_font = style.heading_font or style.font
    parts = [f"<html><body style=\"font-family:'{style.font}'; font-size:{style.size}pt;\">"]
    in_refs = False
    seen_h1 = False
    if style.title_page and title:
        parts.append('<div style="height:200px"></div>')
        parts.append(f'<p align="center" style="font-family:\'{heading_font}\'; font-size:{style.heading_sizes[0] + 4}pt; '
                     f'font-weight:600; color:#{style.heading_color};">{_esc(title)}</p>')
        if author:
            parts.append(f'<p align="center">{_esc(author)}</p>')
        parts.append('<p style="page-break-before:always"></p>')
        seen_h1 = True
    for b in blocks:
        if b.kind == "h":
            counts["h"] += 1
            text = plain(b.runs)
            in_refs = bool(_HEADING_WORDS.fullmatch(text.strip().lower()))
            size = style.heading_sizes[b.level - 1]
            centre = ' align="center"' if (b.level == 1 and (style.heading_center_h1 or (style.center_title and not seen_h1))) else ""
            weight = "600" if style.heading_bold else "400"
            border = "border-bottom:1px solid #bbbbbb;" if (style.rule_under_h2 and b.level == 2) else ""
            top = 4 if style.line_spacing >= 2 else (18 if b.level == 1 else 12)
            # a styled paragraph, not an h tag: Qt scales h tags again on top of
            # the size given, and MLA's 12-point headings came out near 17
            italic = "font-style:italic;" if (b.level == 3 and style.heading_bold and style.heading_center_h1) else ""
            parts.append(f'<p{centre} style="font-family:\'{heading_font}\'; font-size:{size}pt; font-weight:{weight}; {italic}'
                         f'color:#{style.heading_color}; margin-top:{top}px; margin-bottom:4px; {border}">'
                         f'{_html_runs(b.runs)}</p>')
            counts["words"] += len(text.split())
            seen_h1 = seen_h1 or b.level == 1
        elif b.kind == "p":
            hang = in_refs and style.hanging_bibliography
            indent = "margin-left:48px; text-indent:-48px;" if hang else f"text-indent:{style.first_line_indent_in * 96:.0f}px;"
            parts.append(f'<p style="line-height:{line}%; margin-top:0px; margin-bottom:{style.space_after * 1.33:.0f}px; {indent}">'
                         f'{_html_runs(b.runs)}</p>')
            counts["words"] += len(plain(b.runs).split())
        elif b.kind in ("ul", "ol"):
            counts["list"] += 1
            tag = "ul" if b.kind == "ul" else "ol"
            parts.append(f'<{tag} style="margin-top:0px; margin-bottom:6px;">')
            for depth, runs in b.items:
                parts.append(f'<li style="line-height:{line}%; margin-left:{depth * 18}px;">{_html_runs(runs)}</li>')
                counts["words"] += len(plain(runs).split())
            parts.append(f"</{tag}>")
        elif b.kind == "table":
            counts["table"] += 1
            parts.append('<table border="1" cellspacing="0" cellpadding="6" width="100%" style="border-color:#c9c9cf;">')
            for r, row in enumerate(b.rows):
                parts.append("<tr>")
                for c, cell in enumerate(row):
                    align = b.aligns[c] if c < len(b.aligns) else "left"
                    tag = "th" if r == 0 else "td"
                    bg = ' bgcolor="#e9e9ec"' if r == 0 else ""
                    parts.append(f'<{tag}{bg} align="{align}">{_html_runs(cell)}</{tag}>')
                    counts["words"] += len(plain(cell).split())
                parts.append("</tr>")
            parts.append("</table><p></p>")
        elif b.kind == "code":
            counts["code"] += 1
            parts.append('<pre style="font-family:Consolas,monospace; font-size:%.1fpt; background-color:#f2f2f4; '
                         'margin-top:2px; margin-bottom:8px;">%s</pre>' % (max(8.5, style.size - 2), _esc(b.text)))
        elif b.kind == "quote":
            parts.append(f'<blockquote style="margin-left:36px; color:#444444;"><i>{_html_runs(b.runs)}</i></blockquote>')
            counts["words"] += len(plain(b.runs).split())
        elif b.kind == "hr":
            parts.append("<hr/>")
        elif b.kind == "pagebreak":
            parts.append('<p style="page-break-before:always"></p>')
        elif b.kind == "image":
            source = Path(b.text).expanduser()
            if source.is_file():
                counts["image"] += 1
                parts.append(f'<p align="center"><img src="{source.as_uri()}" width="440"/></p>')
                if b.alt:
                    parts.append(f'<p align="center"><i>{_esc(b.alt)}</i></p>')
            else:
                parts.append(f"<p>[image not found: {_esc(b.text)}]</p>")
    parts.append("</body></html>")
    return "\n".join(parts), counts


def write_pdf(path: Path, markdown: str, style: Style, title: str = "", author: str = "") -> dict:
    from PySide6.QtCore import QMarginsF, QRectF, QSizeF, Qt
    from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPainter, QPdfWriter, QTextDocument

    # Mike's own application already exists; on its own (a script, a test) make the
    # full kind, so anything that wants widgets later isn't left beside a bare one.
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    html, counts = _to_html(parse(markdown), style, title, author)
    path.parent.mkdir(parents=True, exist_ok=True)
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4 if style.page_size == "a4" else QPageSize.PageSizeId.Letter))
    margin = style.margin_in * 25.4
    writer.setPageMargins(QMarginsF(margin, margin, margin, margin), QPageLayout.Unit.Millimeter)
    writer.setResolution(300)
    if title:
        writer.setTitle(title)
    writer.setCreator("Mike")
    rect = writer.pageLayout().paintRectPixels(writer.resolution())
    footer = int(0.35 * writer.resolution()) if style.page_numbers == "bottom-center" else 0
    header = int(0.35 * writer.resolution()) if (style.page_numbers == "top-right" or style.header_name) else 0
    text_h = rect.height() - footer - header

    doc = QTextDocument()
    doc.documentLayout().setPaintDevice(writer)
    doc.setDefaultFont(QFont(style.font, int(style.size)))
    # A printed page is dark ink on white whatever the desktop's theme: without
    # this, Qt took the dark-mode text colour and every paragraph came out white.
    doc.setDefaultStyleSheet("body, p, li, td, th, pre, blockquote, span, b, i { color:#141416; }")
    doc.setPageSize(QSizeF(rect.width(), text_h))
    doc.setHtml(html)
    doc.setPageSize(QSizeF(rect.width(), text_h))

    painter = QPainter(writer)
    pages = max(1, doc.pageCount())
    numbering = QFont(style.font, int(style.size))
    last = author.strip().split()[-1] if (style.header_name and author.strip()) else ""
    for i in range(pages):
        painter.save()
        painter.translate(0, header)
        painter.setClipRect(QRectF(0, 0, rect.width(), text_h))
        painter.translate(0, -i * text_h)
        doc.drawContents(painter, QRectF(0, i * text_h, rect.width(), text_h))
        painter.restore()
        painter.setFont(numbering)
        label = f"{last} {i + 1}".strip() if style.header_name else str(i + 1)
        if style.page_numbers == "top-right":
            painter.drawText(QRectF(0, 0, rect.width(), header), int(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter), label)
        elif style.page_numbers == "bottom-center":
            painter.drawText(QRectF(0, header + text_h, rect.width(), footer),
                             int(Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter), label)
        if i < pages - 1:
            writer.newPage()
    painter.end()
    counts["pages"] = pages
    return counts


# ── the front door ───────────────────────────────────────────

def create(path: str, content: str, style_name: str | None = None, title: str = "", author: str = "",
           overwrite: bool = False, toc: bool = False, **override) -> dict:
    """Make a .docx, .pdf, .md or .txt file from Markdown content."""
    from tools.filesystem.path_utils import resolve_path
    file = resolve_path(path)
    suffix = file.suffix.lower()
    if suffix not in (".docx", ".pdf", ".md", ".txt"):
        return {"status": "error", "retry_safe": True,
                "error": f"I make .docx (Word), .pdf, .md and .txt files; '{suffix or path}' isn't one of them. "
                         "Name the file with one of those endings."}
    if not (content or "").strip():
        return {"status": "error", "error": "There's no content to put in the document.", "retry_safe": True}
    if file.is_dir():
        return {"status": "error", "error": f"{file} is a folder.", "retry_safe": True}
    existed = file.exists()
    if existed and not overwrite:
        return {"status": "error", "retry_safe": True,
                "error": f"{file.name} already exists in {file.parent}. Nothing was changed. Use a new name, "
                         "or set overwrite to true if the student wants it replaced."}
    if existed:
        try:
            from brain import revert_store
            revert_store.capture(str(file))
        except Exception:
            logger.debug("Couldn't keep a copy before replacing %s", file, exc_info=True)
    style = style_for(style_name, **override)
    try:
        if suffix == ".docx":
            counts = write_docx(file, content, style, title=title, author=author, toc=toc)
        elif suffix == ".pdf":
            counts = write_pdf(file, content, style, title=title, author=author)
        else:
            file.parent.mkdir(parents=True, exist_ok=True)
            file.write_text(content, encoding="utf-8")
            counts = {"words": len(content.split())}
    except PermissionError:
        return {"status": "error", "retry_safe": False,
                "error": f"{file.name} is open in another program, so it can't be written. Close it and try again."}
    except Exception as exc:
        logger.exception("Making %s failed", file)
        return {"status": "error", "retry_safe": False, "error": f"Couldn't make {file.name}: {exc}"}

    # read it back: what's in the file, not what was asked for
    summary = {"status": "success", "path": str(file), "replaced": existed,
               "style": (style_name or "plain").lower(), "words": counts.get("words", 0)}
    for key, name in (("h", "headings"), ("table", "tables"), ("image", "images"), ("code", "code_blocks")):
        if counts.get(key):
            summary[name] = counts[key]
    detail = ""
    if suffix in (".docx", ".pdf"):
        try:
            from tools.documents import pages
            back = pages.info(str(file))
            if back["count"] < 1:
                raise ValueError("no pages")
            unit = "page" if suffix == ".pdf" else "section"
            summary["pages" if suffix == ".pdf" else "sections"] = back["count"]
            detail = f"{back['count']} {unit}{'s' if back['count'] != 1 else ''}, "
        except Exception as exc:
            return {"status": "error", "retry_safe": False,
                    "error": f"{file.name} was written but doesn't read back properly ({exc}). Don't hand it over."}
    kind = {".docx": "Word document", ".pdf": "PDF", ".md": "Markdown file", ".txt": "text file"}[suffix]
    summary["result"] = f"Made {file.name}, a {kind}: {detail}about {summary['words']:,} words."
    return summary

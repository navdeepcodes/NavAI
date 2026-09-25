"""Put text into a Word document properly — under a heading, keeping the rest.

A .docx is a package, not text, and writing text into one destroys it (that
happened: Mike "finished" a lab report by writing into it as plain text and
the student's document no longer opened). This edits the document as a
document: the section under a heading is replaced, or the heading is added at
the end, and everything else — other sections, formatting, the title — is
left exactly as it was.

Before writing, the previous version is copied aside, so any change can be
undone by hand. After writing, the file is read back and the result reported
from what is actually in it. If Word has the file open it is locked, and the
answer says so rather than failing in some vaguer way.
"""
from __future__ import annotations

import re
import shutil
import time
from pathlib import Path

from brain import mission_checks as checks
from hostplatform import storage
from tools.filesystem.path_utils import resolve_path


def _backup(file: Path) -> Path:
    folder = storage.data_dir() / "backups"
    folder.mkdir(parents=True, exist_ok=True)
    copy = folder / f"{file.stem} before Mike {time.strftime('%Y-%m-%d %H%M%S')}{file.suffix}"
    shutil.copy2(file, copy)
    return copy


def write_section(path: str, heading: str, text: str) -> dict:
    file = resolve_path(path)
    if file.suffix.lower() != ".docx":
        return {"status": "error", "error": (
            f"{file.name} isn't a Word document. This is for .docx files; "
            "for a plain text file use write_file.")}
    if not file.exists():
        return {"status": "error", "error": f"No such file: {file}"}
    name = checks.section_name(heading)
    if not name:
        return {"status": "error", "error": "A heading is needed, e.g. 'Discussion'."}
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", text or "") if p.strip()]
    if not paragraphs:
        return {"status": "error", "error": "There is no text to write."}

    from docx import Document

    try:
        doc = Document(str(file))
    except Exception as exc:
        return {"status": "error", "error": f"Couldn't open {file.name} as a Word document: {exc}"}

    paras = list(doc.paragraphs)
    at = next((i for i, p in enumerate(paras)
               if checks.paragraph_kind(p) == "heading"
               and checks.find_section({checks.section_name(p.text): 0}, name)), None)
    replaced = at is not None
    if replaced:
        anchor = paras[at]
        j = at + 1
        while j < len(paras) and checks.paragraph_kind(paras[j]) != "heading":
            element = paras[j]._p
            element.getparent().remove(element)
            j += 1
        for body in reversed(paragraphs):
            new = doc.add_paragraph(body)
            anchor._p.addnext(new._p)
    else:
        uses_styles = any((p.style.name or "").lower().startswith("heading")
                          for p in paras if p.style is not None)
        if uses_styles:
            doc.add_heading(name, 1)
        else:
            doc.add_paragraph().add_run(name).bold = True
        for body in paragraphs:
            doc.add_paragraph(body)

    backup = _backup(file)
    try:
        doc.save(str(file))
    except PermissionError:
        return {"status": "error", "error": (
            f"{file.name} is open in Word, which locks it, so nothing was written. "
            "Ask the user to save and close it (or paste the text in themselves), then try again.")}

    shape = checks.outline(str(file)) or {}
    found = checks.find_section(shape.get("sections", {}), name)
    words = shape.get("sections", {}).get(found, 0) if found else 0
    if not found or words == 0:
        return {"status": "error", "error": (
            f"Wrote to {file.name}, but reading it back the “{name}” section isn't there. "
            f"The previous version is saved at {backup}.")}
    return {
        "status": "success",
        "result": (f"{'Replaced' if replaced else 'Added'} the “{found}” section in {file.name}: "
                   f"{words} words, read back from the file. The previous version is saved as "
                   f"“{backup.name}” in Mike's backups folder."),
        "words": words,
        "backup": str(backup),
    }

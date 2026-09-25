"""Writing into a Word document edits it as a document — never as text.

Measured: asked to help finish a lab report, Mike wrote the answer into the
.docx as plain text and the student's report no longer opened.
"""
from pathlib import Path

import pytest
from docx import Document

from brain import mission_checks as checks
from tools.filesystem.document_writer import write_section


def _report(path: Path) -> Path:
    d = Document()
    d.add_heading("Lab 3", 0)
    d.add_heading("Aim", 1)
    d.add_paragraph("To measure g with a pendulum.")
    d.add_heading("Results", 1)
    d.add_paragraph("old results text")
    d.add_heading("Conclusion", 1)
    d.add_paragraph("It worked.")
    d.save(str(path))
    return path


def test_a_missing_section_is_added_and_the_rest_is_untouched(tmp_path):
    report = _report(tmp_path / "r.docx")
    text = "Our g was close to 9.81.\n\nReaction time mattered most."
    r = write_section(str(report), "Discussion", text)
    assert r["status"] == "success" and "Added" in r["result"]
    shape = checks.outline(str(report))
    assert shape["sections"]["Discussion"] == checks.count_words(text)
    assert shape["sections"]["Aim"] == 6 and shape["sections"]["Conclusion"] == 2
    assert Document(str(report)).paragraphs[0].text == "Lab 3"        # title kept


def test_an_existing_section_is_replaced_in_place(tmp_path):
    report = _report(tmp_path / "r.docx")
    text = "g = 9.80 m/s2 from the gradient of T squared against L."
    r = write_section(str(report), "Results", text)
    assert r["status"] == "success" and "Replaced" in r["result"]
    shape = checks.outline(str(report))
    assert shape["order"] == ["Aim", "Results", "Conclusion"]
    assert "old results" not in " ".join(p.text for p in Document(str(report)).paragraphs)
    assert shape["sections"]["Results"] == checks.count_words(text)


def test_the_previous_version_is_kept(tmp_path):
    report = _report(tmp_path / "r.docx")
    r = write_section(str(report), "Results", "new")
    backup = Path(r["backup"])
    assert backup.exists() and "old results text" in " ".join(p.text for p in Document(str(backup)).paragraphs)


def test_a_file_open_in_word_is_reported_not_half_written(tmp_path, monkeypatch):
    report = _report(tmp_path / "r.docx")
    before = report.read_bytes()

    def locked(self, *_a, **_k):
        raise PermissionError("locked")

    monkeypatch.setattr("docx.document.Document.save", locked)
    r = write_section(str(report), "Discussion", "text")
    assert r["status"] == "error" and "open in Word" in r["error"]
    assert report.read_bytes() == before


def test_only_word_documents(tmp_path):
    notes = tmp_path / "n.txt"
    notes.write_text("x", encoding="utf-8")
    assert write_section(str(notes), "Aim", "text")["status"] == "error"


def test_a_write_waits_for_the_users_ok_and_no_leaves_the_file_alone(tmp_path, monkeypatch):
    """Whether to write is the model's call and the user's consent -- not a
    guess from the wording. The prompt shows what would be replaced."""
    from tests.test_unacted_promise import _Scripted
    from brain.core_runtime import CoreRuntime
    from docx import Document

    report = tmp_path / "r.docx"
    d = Document(); d.add_heading("Results", 1); d.add_paragraph("word " * 60); d.save(str(report))
    before = report.read_bytes()
    brain = _Scripted(("write_document_section", {"path": str(report), "heading": "Results", "text": "new"}),
                      "Okay, I've left it.")
    rt = CoreRuntime(); rt._brain = brain; rt._capabilities = brain.capabilities()
    asked = []
    events = list(rt.process_streaming("what should I write in the results?",
                                       confirm_callback=lambda detail: asked.append(detail) or False))
    assert len(asked) == 1 and "Replace the “Results” section (60 words now)" in asked[0]
    assert any(k == "tool_end" and "denied" in str(p) for k, p in events)
    assert report.read_bytes() == before


def test_writing_needs_the_users_ok():
    from brain.core_tools import needs_confirmation
    assert needs_confirmation("write_document_section", {"path": "r.docx", "heading": "Aim", "text": "x"})

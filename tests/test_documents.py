"""Documents and PDFs for a student: long reading with page numbers, search,
OCR of scans, making Word/PDF/slide files, and PDF tools.

The document tests use real files made in a temp folder (PDFs written by Mike's
own PDF maker, Word and PowerPoint through python-docx/python-pptx, scans as
pictures), so they check what a student's files actually do. OCR ones run only
where Windows' engine is present.
"""
from __future__ import annotations

import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from tests import _isolate  # noqa: F401,E402

from tools.documents import create, ocr, pages, pdf_ops, slides  # noqa: E402
from tools.documents.tool import TOOLS  # noqa: E402

ESSAY = """# Printing

Gutenberg's press made **identical** pages possible.

## Standardization

Before print every copy differed. After print, readers could cite the same page.

- one
- two

| Era | Copies |
|-----|-------:|
| Manuscript | 100 |
| Print | 10,000 |

1. first
2. second

```python
print("hi")
```

## Works Cited

Eisenstein, Elizabeth. *The Printing Press as an Agent of Change*. 1979.
"""


def _long_markdown(sections: int = 14) -> str:
    out = ["# The Handbook"]
    for i in range(1, sections + 1):
        out.append(f"## Chapter {i}")
        out.append(("Chapter %d discusses topic%d in detail. " % (i, i)) * 60)
    return "\n\n".join(out)


@pytest.fixture
def work(tmp_path, monkeypatch):
    from hostplatform import storage
    monkeypatch.setattr(storage, "data_dir", lambda: tmp_path / "data")
    return tmp_path


def _make(work, name, md=ESSAY, style="plain", **kw):
    path = work / name
    result = create.create(str(path), md, style, overwrite=True, **kw)
    assert result["status"] == "success", result
    return path


# ── page ranges ──────────────────────────────────────────────

def test_page_specs_are_read_the_way_a_student_writes_them():
    assert pages.parse_pages("3", 10) == [3]
    assert pages.parse_pages("3-5", 10) == [3, 4, 5]
    assert pages.parse_pages("1,4,9-10", 10) == [1, 4, 9, 10]
    assert pages.parse_pages("8-", 10) == [8, 9, 10]
    assert pages.parse_pages("last", 10) == [10]
    assert pages.parse_pages(" 2 to 4 ", 10) == [2, 3, 4]
    assert pages.parse_pages(None, 10) == []
    assert pages.ranges([1, 2, 3, 7, 9, 10]) == "1-3,7,9-10"
    for bad in ("11", "0", "5-3", "abc"):
        with pytest.raises(ValueError):
            pages.parse_pages(bad, 10)


# ── reading and searching ────────────────────────────────────

def test_a_long_word_file_is_read_in_pieces_and_says_what_comes_next(work):
    file = _make(work, "handbook.docx", _long_markdown())
    info = pages.info(str(file))
    assert info["count"] >= 10 and info["outline"][0].startswith("The Handbook")
    first = pages.read(str(file))
    assert first["status"] == "success" and first["shown"].startswith("1") and first["next"]
    assert len(first["text"]) <= pages.READ_CHARS + 200
    later = pages.read(str(file), "last")
    assert "Chapter 14" in later["text"] and "next" not in later


def test_searching_says_which_page_and_matches_across_line_breaks(work):
    long_pdf = _make(work, "handbook.pdf", _long_markdown(8))
    found = pages.search("topic6 in detail", str(long_pdf))
    assert found["status"] == "success" and found["hits"], found
    assert all("topic6" in " ".join(h["snippets"]) for h in found["hits"])
    assert found["hits"][0]["label"].startswith("Page ")


def test_a_folder_is_searched_with_everything_in_it(work):
    _make(work, "a.docx", "# A\n\nThe mitochondria is the powerhouse of the cell.\n")
    _make(work, "b.pdf", "# B\n\nNothing relevant here at all.\n")
    (work / "c.md").write_text("Notes: the mitochondria matter for exams.", encoding="utf-8")
    found = pages.search("mitochondria", str(work))
    assert {h["file"] for h in found["hits"]} == {"a.docx", "c.md"}


def test_words_are_found_when_the_exact_phrase_is_not(work):
    _make(work, "n.docx", "# N\n\nEnergy is stored in glucose during photosynthesis in the leaf.\n")
    found = pages.search("photosynthesis glucose energy", str(work))
    assert found["matched_on"] == "words" and found["hits"]


def test_a_missing_or_wrong_page_is_said_plainly(work):
    file = _make(work, "e.docx")
    bad = pages.read(str(file), "99")
    assert bad["status"] == "error" and "isn't there" in bad["error"] and bad["retry_safe"]
    with pytest.raises(FileNotFoundError):
        pages.open_doc(str(work / "nope.pdf"))


# ── scans and pictures ───────────────────────────────────────

needs_ocr = pytest.mark.skipif(not ocr.available(), reason="Windows OCR isn't available here")


def _scanned_pdf(work):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1000, 320), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(r"C:\Windows\Fonts\times.ttf", 36)
    draw.text((30, 40), "Lecture four covers photosynthesis", font=font, fill="black")
    draw.text((30, 110), "and the Calvin cycle in detail.", font=font, fill="black")
    picture = work / "scan.png"
    img.save(picture)
    pdf = work / "scan.pdf"
    img.save(pdf, "PDF", resolution=110.0)
    return picture, pdf


@needs_ocr
def test_a_scanned_pdf_is_read_by_ocr_and_says_so(work):
    _picture, pdf = _scanned_pdf(work)
    assert pages.info(str(pdf))["scanned"] is True
    got = pages.read(str(pdf))
    assert "photosynthesis" in got["text"].lower() and any("OCR" in n for n in got["notes"])
    again = pages.read(str(pdf))                       # remembered: no second OCR
    assert "photosynthesis" in again["text"].lower()
    assert any((work / "data" / "doc_cache").glob("*.json"))


@needs_ocr
def test_a_photo_of_notes_is_read_and_searchable(work):
    picture, _pdf = _scanned_pdf(work)
    assert "Calvin" in pages.read(str(picture))["text"]
    assert pages.search("calvin cycle", str(work))["hits"]


# ── making documents ─────────────────────────────────────────

def test_an_mla_essay_has_the_formatting_mla_asks_for(work):
    from docx import Document
    file = _make(work, "mla.docx", style="mla", author="Thrisha Reddy", title="Printing")
    doc = Document(str(file))
    assert doc.styles["Normal"].font.name == "Times New Roman" and doc.styles["Normal"].font.size.pt == 12
    body = [p for p in doc.paragraphs if p.text.startswith("Gutenberg")][0]
    assert body.paragraph_format.line_spacing == 2.0 and abs(body.paragraph_format.first_line_indent.inches - 0.5) < 0.01
    header = doc.sections[0].header.paragraphs[0]
    assert header.text.startswith("Reddy") and "PAGE" in header._p.xml
    cited = [p for p in doc.paragraphs if p.text.startswith("Eisenstein")][0]
    assert cited.paragraph_format.left_indent.inches == 0.5 and cited.paragraph_format.first_line_indent.inches == -0.5


def test_apa_gets_a_title_page_and_a_report_a_contents_list(work):
    apa = _make(work, "apa.pdf", style="apa", author="A. Student", title="Printing and Power")
    first = pages.read(str(apa), "1")["text"]
    assert "Printing and Power" in first and "A. Student" in first and "Gutenberg" not in first
    assert pages.info(str(apa))["count"] >= 2
    report = _make(work, "report.docx", style="report", title="Report", toc=True)
    from docx import Document
    assert "TOC" in Document(str(report)).element.xml


def test_lists_tables_code_and_emphasis_survive_in_word_and_pdf(work):
    from docx import Document
    doc = Document(str(_make(work, "all.docx")))
    assert len(doc.tables) == 1 and doc.tables[0].cell(1, 1).text == "100"
    assert any(r.bold and r.text == "identical" for p in doc.paragraphs for r in p.runs)
    assert any(p.text.startswith("•") for p in doc.paragraphs) and any(p.text.startswith("2.") for p in doc.paragraphs)
    text = pages.read(str(_make(work, "all.pdf")))["text"]
    for needle in ("Manuscript", "Copies", "print(\"hi\")", "second"):
        assert needle in text


def test_a_document_that_exists_is_never_replaced_without_being_told(work):
    file = _make(work, "keep.docx")
    again = create.create(str(file), "# other", "plain")
    assert again["status"] == "error" and "already exists" in again["error"] and again["retry_safe"]
    assert create.create(str(file), "# replaced", "plain", overwrite=True)["replaced"] is True


def test_the_formats_and_content_are_checked(work):
    assert "isn't one of them" in create.create(str(work / "x.exe"), "text")["error"]
    assert "no content" in create.create(str(work / "x.docx"), "   ")["error"]
    result = create.create(str(work / "notes.md"), "# hi\n")
    assert result["status"] == "success" and (work / "notes.md").read_text(encoding="utf-8") == "# hi\n"


def test_a_resume_is_one_tight_page(work):
    md = "# Priya Nair\n\npriya@example.com | Chennai\n\n## Education\n\n**B.Tech, Computer Science** — Anna University, 2027\n\n## Projects\n\n- Built a study planner with Flask\n- Led a team of 3 at a 24-hour hackathon\n"
    file = _make(work, "resume.pdf", md, style="resume")
    assert pages.info(str(file))["count"] == 1


def test_a_deck_is_a_real_presentation_with_titles_and_notes(work):
    from pptx import Presentation
    result = slides.create_deck(str(work / "deck.pptx"), "Cell Biology", [
        {"title": "Cell Biology", "subtitle": "BIO 101"},
        {"title": "Organelles", "bullets": ["Nucleus", "Mitochondria", ["Makes ATP"]], "notes": "Say: energy first."},
        {"title": "Comparison", "table": [["Part", "Job"], ["Nucleus", "DNA"]]},
        {"title": "Quote", "quote": "Life is chemistry.", "attribution": "Someone"},
        {"title": "Two sides", "left": ["a"], "right": ["b"], "left_title": "L", "right_title": "R"},
        {"title": "Next up"},
        {"title": "Missing picture", "image": str(work / "nope.png")},
    ], theme="graphite")
    assert result["status"] == "success" and result["slides"] == 7 and result["problems"]
    deck = Presentation(str(work / "deck.pptx"))
    assert deck.slides[1].shapes.title.text_frame.text == "Organelles"
    assert deck.slides[1].notes_slide.notes_text_frame.text == "Say: energy first."
    outline = pages.info(str(work / "deck.pptx"))["outline"]
    assert outline[1].startswith("Organelles") and pages.read(str(work / "deck.pptx"), "2")["text"].count("Mitochondria") == 1
    exists = slides.create_deck(str(work / "deck.pptx"), "x", [{"title": "y"}])
    assert exists["status"] == "error" and "already exists" in exists["error"]
    assert slides.create_deck(str(work / "z.txt"), "x", [{"title": "y"}])["status"] == "error"


# ── PDF tools ────────────────────────────────────────────────

def _pdf(work, name, words, sections=6):
    return _make(work, name, "\n\n".join(f"## {words} {i}\n\n" + (f"{words} page {i} text. " * 80) + "\n\n\\pagebreak"
                                          for i in range(1, sections + 1)))


def test_merge_extract_delete_rotate_reorder_split_never_touch_the_original(work):
    a, b = _pdf(work, "a.pdf", "Alpha", 3), _pdf(work, "b.pdf", "Beta", 2)
    original = a.read_bytes()
    merged = pdf_ops.run("merge", paths=[str(a), str(b)])
    assert merged["status"] == "success" and merged["pages"] == pages.info(str(a))["count"] + pages.info(str(b))["count"]
    made = pdf_ops.run("extract", path=merged["path"], pages="2-3")
    assert made["pages"] == 2
    assert pdf_ops.run("delete_pages", path=merged["path"], pages="1")["pages"] == merged["pages"] - 1
    assert pdf_ops.run("rotate", path=str(a), degrees=90, pages="1")["status"] == "success"
    order = pdf_ops.run("reorder", path=str(a), order="3,1")
    assert order["pages"] == 2 and "Alpha" in pages.read(order["path"], "1")["text"]
    parts = pdf_ops.run("split", path=merged["path"], every=2)
    assert parts["count"] >= 2 and all(f["pages"] <= 2 for f in parts["files"])
    assert a.read_bytes() == original, "the student's own file is untouched"
    assert not pdf_ops.run("extract", path=str(a), pages="99")["status"] == "success"


def test_results_never_overwrite_an_existing_file(work):
    a = _pdf(work, "a.pdf", "Alpha", 3)
    one = pdf_ops.run("extract", path=str(a), pages="1")
    two = pdf_ops.run("extract", path=str(a), pages="1")
    assert one["path"] != two["path"] and (work / "a (pages 1).pdf").exists()


def test_photos_become_one_pdf_and_a_pdf_can_be_made_smaller(work):
    from PIL import Image
    shots = []
    for i, colour in enumerate(("white", "lightgrey", "white")):
        p = work / f"p{i}.jpg"
        Image.new("RGB", (1200, 1600), colour).save(p, quality=95)
        shots.append(str(p))
    made = pdf_ops.run("from_images", paths=shots)
    assert made["pages"] == 3
    small = pdf_ops.run("compress", path=made["path"])
    assert small["status"] == "success" and small["size_kb"] <= small["before_kb"]


def _form_pdf(work):
    """A tiny PDF with one text field and one checkbox, built by hand."""
    from pypdf import PdfWriter
    from pypdf.generic import (ArrayObject, DictionaryObject, NameObject, NumberObject, TextStringObject)
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    fields = ArrayObject()
    for name, kind, rect in (("full_name", "/Tx", [20, 240, 200, 260]), ("agree", "/Btn", [20, 200, 40, 220])):
        widget = DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"),
            NameObject("/FT"): NameObject(kind), NameObject("/T"): TextStringObject(name),
            NameObject("/Rect"): ArrayObject([NumberObject(v) for v in rect]), NameObject("/F"): NumberObject(4)})
        ref = writer._add_object(widget)
        fields.append(ref)
        if NameObject("/Annots") not in page:
            page[NameObject("/Annots")] = ArrayObject()
        page[NameObject("/Annots")].append(ref)
    writer._root_object[NameObject("/AcroForm")] = DictionaryObject({NameObject("/Fields"): fields})
    path = work / "form.pdf"
    with open(path, "wb") as handle:
        writer.write(handle)
    return path


def test_a_form_is_listed_filled_and_checked(work):
    form = _form_pdf(work)
    listed = pdf_ops.run("fields", path=str(form))
    assert {f["name"] for f in listed["fields"]} == {"full_name", "agree"}
    filled = pdf_ops.run("fill", path=str(form), values={"full_name": "Priya Nair", "Agree": True, "nope": "x"})
    assert filled["status"] == "success" and filled["filled"] == 2 and "nope" in filled["not_found"][0]
    values = {f["name"]: f["value"] for f in pdf_ops.run("fields", path=filled["path"])["fields"]}
    assert values["full_name"] == "Priya Nair"
    flat = pdf_ops.run("fields", path=str(_pdf(work, "flat.pdf", "Flat", 1)))
    assert flat["count"] == 0 and "flat" in flat["result"]
    assert pdf_ops.run("fill", path=str(_pdf(work, "flat2.pdf", "Flat", 1)), values={"a": "b"})["status"] == "error"


def test_conversions_between_word_pdf_and_text(work):
    docx = _make(work, "paper.docx", ESSAY)
    md = pdf_ops.run("convert", path=str(docx), to="md")
    text = (work / os.path.basename(md["path"])).read_text(encoding="utf-8")
    assert "## Standardization" in text and "| Manuscript | 100 |" in text
    pdf = _make(work, "paper.pdf", ESSAY)
    word = pdf_ops.run("convert", path=str(pdf), to="docx")
    assert word["status"] == "success" and "Only the text" in word["result"]
    assert "Gutenberg" in pages.read(word["path"])["text"]
    assert pdf_ops.run("convert", path=str(pdf), to="pdf")["status"] == "error"
    assert pdf_ops.run("convert", path=str(docx), to="xlsx")["status"] == "error"


@pytest.mark.skipif(not pdf_ops._office_available("Word.Application"), reason="Word isn't installed")
def test_word_makes_the_pdf_when_it_is_installed(work):
    docx = _make(work, "for_word.docx", ESSAY)
    done = pdf_ops.run("convert", path=str(docx), to="pdf")
    assert done["status"] == "success" and "Gutenberg" in pages.read(done["path"])["text"]


# ── through Mike's runtime ───────────────────────────────────

def test_every_document_tool_runs_through_the_runtime_and_can_be_switched_off(work, monkeypatch):
    from brain import permissions
    from brain.core_runtime import CoreRuntime
    from brain.core_tools import needs_confirmation
    runtime = CoreRuntime.__new__(CoreRuntime)
    made = runtime._execute_tool("create_document", {"path": str(work / "r.docx"), "content": ESSAY, "style": "mla"})
    assert made["status"] == "success" and "Made r.docx" in made["result"]
    read = runtime._execute_tool("read_document", {"path": str(work / "r.docx")})
    assert read["status"] == "success" and "Section" in read["result"] and "Gutenberg" in read["result"]
    assert runtime._execute_tool("document_info", {"path": str(work / "r.docx")})["count"] >= 1
    assert runtime._execute_tool("search_document", {"query": "citation pages", "path": str(work)})["status"] == "success"
    assert runtime._execute_tool("pdf_edit", {"action": "info", "path": str(_make(work, "i.pdf"))})["status"] == "success"
    assert runtime._execute_tool("create_presentation", {"path": str(work / "d.pptx"), "slides": [{"title": "Hi"}]})["status"] == "success"
    # a new file needs no approval; replacing one does
    assert not needs_confirmation("create_document", {"path": str(work / "new.docx"), "content": "x"})
    assert not needs_confirmation("create_document", {"path": str(work / "nothere.docx"), "overwrite": True})
    assert needs_confirmation("create_document", {"path": str(work / "r.docx"), "overwrite": True})
    monkeypatch.setattr(permissions, "_pref_set", lambda name: {"documents"} if name == "abilities_off" else set())
    off = runtime._execute_tool("create_document", {"path": str(work / "q.docx"), "content": "x"})
    assert off["status"] == "error" and "Work with documents" in off["error"] and not (work / "q.docx").exists()


# ── attached files ───────────────────────────────────────────

def test_an_attached_file_arrives_with_its_path_so_mike_can_do_things_with_it(work):
    from tools.documents import attach
    a, b = _pdf(work, "week1.pdf", "Alpha", 3), _pdf(work, "week2.pdf", "Beta", 2)
    text = attach.describe_all([str(a), str(b)])
    for path in (a, b):
        assert f"path: {path}" in text, "the tools need the full path"
    assert "PDF, " in text and "pages" in text and "Alpha page 1" in text
    assert "Read the rest with read_document" in text


def test_a_photo_of_notes_is_read_by_ocr_and_not_described_by_the_slow_model(work):
    from tools.documents import attach
    asked = []
    picture, _pdf_ = _scanned_pdf(work) if ocr.available() else (None, None)
    if picture is None:
        pytest.skip("Windows OCR isn't available here")
    text = attach.describe(str(picture), describe_image=lambda p: asked.append(p) or "a diagram")
    assert "photosynthesis" in text.lower() and "read by OCR" in text
    assert asked == [], "a picture with text needs no picture model"


def test_a_picture_with_no_text_is_described_and_a_missing_file_is_said(work):
    from PIL import Image
    from tools.documents import attach
    blank = work / "blank.png"
    Image.new("RGB", (300, 200), "white").save(blank)
    seen = []
    text = attach.describe(str(blank), describe_image=lambda p: seen.append(p) or "a plain white square")
    assert seen == [str(blank)] and "a plain white square" in text and "300x200" in text
    gone = attach.describe(str(work / "gone.pdf"))
    assert "path:" in gone and "isn't there" in gone


def test_the_preview_budget_is_shared_between_attachments(work):
    from tools.documents import attach
    files = [str(_pdf(work, f"f{i}.pdf", f"Name{i}", 3)) for i in range(8)]
    text = attach.describe_all(files)
    assert len(text) < attach.TOTAL_PREVIEW + 8 * 700

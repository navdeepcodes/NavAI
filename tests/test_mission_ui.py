"""The mission in view: the bar above the composer, the corner line, and the
tracker that keeps both true from the document without the model."""
import sys
import time

import pytest

from brain import mission_store as ms


def _app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication(sys.argv)


def _pump(seconds, until=None):
    app = _app()
    end = time.time() + seconds
    while time.time() < end:
        app.processEvents()
        if until and until():
            return True
        time.sleep(0.02)
    return bool(until and until())


@pytest.fixture(autouse=True)
def _clean():
    m = ms.active()
    if m:
        ms.finish(m["id"], "dropped")
    yield
    m = ms.active()
    if m:
        ms.finish(m["id"], "dropped")


M = {"id": 7, "goal": "Finish lab report", "deadline": "tonight", "blocker": "", "status": "active",
     "steps": [
         {"title": "Read the brief", "section": "", "target": 0, "status": "done", "evidence": "you said so"},
         {"title": "Write the Results", "section": "Results", "target": 150, "status": "done",
          "evidence": "160/150 words in Results (r.docx)"},
         {"title": "Write the Discussion", "section": "Discussion", "target": 150, "status": "todo",
          "evidence": "40/150 words in Discussion (r.docx)"}],
     "files": []}


def test_the_bar_shows_the_mission_and_says_when_a_step_completes():
    _app()
    from ui.workspace.mission_bar import MissionBar

    bar = MissionBar()
    bar.set_mission(M, ["Write the Results"])
    assert bar.isVisible() is False or bar.isHidden() is False
    assert bar._flash.startswith("Write the Results — done")
    collapsed = bar.height()
    bar._expanded = True
    bar.set_mission(M)
    assert bar.height() > collapsed                     # every step, with its evidence
    bar.grab()                                          # paints without error


def test_stopping_tracking_takes_two_clicks():
    _app()
    from PySide6.QtCore import QPointF, Qt
    from PySide6.QtGui import QMouseEvent
    from PySide6.QtCore import QEvent
    from ui.workspace.mission_bar import MissionBar

    bar = MissionBar()
    bar.resize(700, 10)
    bar._expanded = True
    bar.set_mission(M)
    bar.resize(700, bar.height())
    dropped = []
    bar.drop_requested.connect(lambda: dropped.append(1))
    at = bar._foot_rect().center()

    def click():
        ev = QMouseEvent(QEvent.MouseButtonRelease, QPointF(at), QPointF(at),
                         Qt.LeftButton, Qt.NoButton, Qt.NoModifier)
        bar.mouseReleaseEvent(ev)

    click()
    assert dropped == [] and bar._armed
    click()
    assert dropped == [1]


def test_a_finished_mission_is_shown_then_settles_away(monkeypatch):
    _app()
    import ui.workspace.mission_bar as mb

    monkeypatch.setattr(mb, "SETTLE_MS", 50)
    bar = mb.MissionBar()
    done = dict(M, status="done", steps=[dict(s, status="done") for s in M["steps"]])
    bar.set_mission(done)
    assert not bar.isHidden()
    assert _pump(1.0, until=bar.isHidden)


def test_the_corner_carries_the_mission_in_one_line():
    _app()
    from ui.workspace.corner import CornerPresence

    corner = CornerPresence()
    corner.set_mission(M, ["Write the Results"])
    assert not corner._mission.isHidden()
    corner.set_mission(None)
    assert corner._mission.isHidden()


def test_the_tracker_follows_the_document_without_the_model(tmp_path):
    _app()
    from docx import Document
    from ui.controller.mission_tracker import MissionTracker

    report = tmp_path / "r.docx"
    d = Document(); d.add_heading("Results", 1); d.add_paragraph("short"); d.save(str(report))
    brief = tmp_path / "brief.txt"
    brief.write_text("Results (about 60 words)\nDiscussion (about 60 words)\n", encoding="utf-8")
    ms.start("Lab report", ["Read the brief", "Write the Results", "Write the Discussion"],
             files=[str(report)], brief=[str(brief)])

    seen = []
    tracker = MissionTracker()
    tracker.changed.connect(lambda m, nd, ch: seen.append((m, nd)))
    tracker._timer.setInterval(200)
    tracker.start()
    assert _pump(5, until=lambda: bool(seen)), "the tracker reports the mission"

    time.sleep(1.1)                                     # a distinct modification time
    d = Document(str(report)); d.add_paragraph("word " * 80); d.save(str(report))
    assert _pump(8, until=lambda: any("Write the Results" in nd for _, nd in seen)), \
        "saving the document ticks the step, with no message sent"
    tracker.stop()


def test_esc_on_an_approval_declines_that_action_only():
    """The card says "Don't do this (Esc)"; Esc used to cancel the whole
    turn instead, so Mike stopped rather than carrying on without it."""
    _app()
    from PySide6.QtCore import QEvent, Qt
    from PySide6.QtGui import QKeyEvent
    from ui.app import MikeWindow

    declined, cancelled = [], []
    win = MikeWindow.__new__(MikeWindow)

    class _Page:
        def __init__(self):
            from ui.workspace.chat_page import ConfirmCard
            self.confirm = ConfirmCard()
            self.confirm.denied.connect(lambda: declined.append(True))

        def showing_overlay(self):
            return False

    class _Controller:
        def cancel_active(self):
            cancelled.append(True)

    win.page, win.controller = _Page(), _Controller()
    win.page.confirm.ask("Add a “Discussion” section in report.docx")
    win.page.confirm.show()
    MikeWindow.keyPressEvent(win, QKeyEvent(QEvent.KeyPress, Qt.Key_Escape, Qt.NoModifier))
    assert declined == [True] and cancelled == []
    win.page.confirm.hide()
    MikeWindow.keyPressEvent(win, QKeyEvent(QEvent.KeyPress, Qt.Key_Escape, Qt.NoModifier))
    assert cancelled == [True]                          # no card: Esc still stops Mike

"""Missions: what the user is trying to get done, with progress that is true.

The central promise is that a mission never says more than the work shows. A
step about a section of a document is ticked by the document, not by the
model, and un-ticked if the section is deleted; finishing is refused while
the document still shows gaps.
"""
import pytest

from brain import mission_checks as checks
from brain import mission_store as ms


@pytest.fixture(autouse=True)
def _no_active_mission():
    m = ms.active()
    if m:
        ms.finish(m["id"], "dropped")
    yield
    m = ms.active()
    if m:
        ms.finish(m["id"], "dropped")


def _docx(path, sections, bold_headings=False, title="Lab 3: Simple Pendulum"):
    from docx import Document

    doc = Document()
    doc.add_heading(title, 0)
    for heading, words in sections:
        if bold_headings:
            doc.add_paragraph().add_run(heading).bold = True
        else:
            doc.add_heading(heading, 1)
        if words:
            doc.add_paragraph(" ".join(["word"] * words))
    doc.save(str(path))
    return path


def _brief(path):
    path.write_text(
        "PHY101 Lab 3 - Simple Pendulum\n"
        "Aim (about 30 words)\n"
        "Method\n"
        "Results (about 150 words)\n"
        "Discussion: 150-200 words\n"
        "Conclusion (80 words)\n", encoding="utf-8")
    return path


# ── reading the work ─────────────────────────────────────

def test_outline_reads_word_headings_and_words_under_each(tmp_path):
    doc = _docx(tmp_path / "report.docx", [("Aim", 40), ("2. Results:", 120), ("Discussion", 0)])
    shape = checks.outline(str(doc))
    assert shape["sections"] == {"Aim": 40, "Results": 120, "Discussion": 0}
    assert shape["words"] == 160


def test_bold_lines_count_as_headings_because_that_is_how_students_write(tmp_path):
    doc = _docx(tmp_path / "report.docx", [("Results", 90), ("Discussion", 10)], bold_headings=True)
    assert checks.outline(str(doc))["sections"] == {"Results": 90, "Discussion": 10}


def test_markdown_and_plain_text_headings(tmp_path):
    md = tmp_path / "notes.md"
    md.write_text("# Aim\nfind g\n\n## Results\n" + "x " * 30 + "\nMethod:\nswing it\n", encoding="utf-8")
    assert checks.outline(str(md))["sections"] == {"Aim": 2, "Results": 30, "Method": 2}


def test_targets_come_from_the_brief(tmp_path):
    brief = _brief(tmp_path / "brief.txt").read_text(encoding="utf-8")
    assert checks.target_words(brief, "Results") == 150
    assert checks.target_words(brief, "Discussion") == 150
    assert checks.target_words(brief, "Method") is None
    assert {"Aim", "Method", "Results", "Discussion", "Conclusion"} <= set(checks.brief_headings(brief))


def test_steps_link_to_the_section_they_name():
    heads = ["Aim", "Results", "Discussion"]
    assert checks.link_step("Write the Results section", heads) == "Results"
    assert checks.link_step("Draft discussion of errors", heads) == "Discussion"
    assert checks.link_step("Proofread everything", heads) is None


# ── the mission itself ───────────────────────────────────

def _lab_mission(tmp_path, results_words=0):
    report = _docx(tmp_path / "lab3_report.docx",
                   [("Aim", 40), ("Method", 90), ("Results", results_words)])
    brief = _brief(tmp_path / "lab3_brief.txt")
    return ms.start("Finish physics lab report", [
        "Read the brief", "Write the Aim", "Write the Results section", "Write the Discussion",
        "Write the Conclusion", "Proofread and submit"],
        files=[str(report)], brief=[str(brief)], deadline="tonight"), report


def test_a_mission_links_steps_to_sections_and_checks_them_from_the_file(tmp_path):
    m, _ = _lab_mission(tmp_path, results_words=40)
    steps = {s["title"]: s for s in m["steps"]}
    assert steps["Write the Results section"]["section"] == "Results"
    assert steps["Write the Results section"]["target"] == 150
    assert steps["Write the Results section"]["status"] == "todo"
    assert "40/150" in steps["Write the Results section"]["evidence"]
    assert steps["Write the Discussion"]["evidence"].startswith("no ")
    assert {f["role"] for f in m["files"]} == {"work", "brief"}
    # Aim is already written in the report: credit from the first moment
    assert steps["Write the Aim"]["status"] == "done"


def test_the_models_plan_is_kept_as_it_gave_it(tmp_path):
    """The model plans; the harness only adds evidence. A step for a section
    the brief never mentions stays, unchecked, until a heading shows it."""
    report = _docx(tmp_path / "lab3_report.docx", [("Aim", 40), ("Method", 90)])
    brief = _brief(tmp_path / "lab3_brief.txt")
    plan = ["Read the lab brief", "Write the Introduction section", "Describe the method",
            "Write conclusion", "Proofread and submit"]
    m = ms.start("Finish lab report", plan, files=[str(report)], brief=[str(brief)])
    assert [s["title"] for s in m["steps"]] == plan
    steps = {s["title"]: s for s in m["steps"]}
    assert steps["Describe the method"]["section"] == "Method"
    assert steps["Describe the method"]["status"] == "done"
    assert not steps["Write the Introduction section"]["section"]


def test_a_heading_added_later_links_the_step_that_names_it(tmp_path):
    report = _docx(tmp_path / "essay.docx", [("Aim", 40)])
    m = ms.start("Finish essay", ["Write the Introduction", "Proofread"], files=[str(report)])
    assert not ms.get(m["id"])["steps"][0]["section"]
    _docx(report, [("Aim", 40), ("Introduction", 70)])
    step = ms.evaluate(m["id"])["mission"]["steps"][0]
    assert step["section"] == "Introduction" and step["status"] == "done"


def test_a_file_nobody_writes_in_is_read_as_the_brief(tmp_path):
    report = _docx(tmp_path / "report.docx", [("Aim", 40)])
    slides = tmp_path / "task.pptx"
    from pptx import Presentation
    Presentation().save(str(slides))
    m = ms.start("Finish report", ["Write the Aim"], files=[str(report), str(slides)])
    from pathlib import Path
    roles = {Path(f["path"]).name: f["role"] for f in m["files"]}
    assert roles == {"report.docx": "work", "task.pptx": "brief"}


def test_writing_the_section_ticks_the_step_and_deleting_it_unticks(tmp_path):
    m, report = _lab_mission(tmp_path, results_words=40)
    _docx(report, [("Aim", 40), ("Method", 90), ("Results", 160)])
    r = ms.evaluate(m["id"])
    assert r["newly_done"] == ["Write the Results section"]
    assert "Results: 40 → 160 words" in r["changes"]
    _docx(report, [("Aim", 40), ("Method", 90)])
    m2 = ms.evaluate(m["id"])["mission"]
    assert next(s for s in m2["steps"] if s["section"] == "Results")["status"] == "todo"


def test_the_model_cannot_tick_a_section_that_is_not_written(tmp_path):
    m, _ = _lab_mission(tmp_path, results_words=40)
    reply = ms.set_step(m["id"], "Results", "done")
    assert reply.startswith("Not marked done")
    assert next(s for s in ms.get(m["id"])["steps"] if s["section"] == "Results")["status"] == "todo"


def test_steps_no_file_can_show_are_marked_as_said(tmp_path):
    m, _ = _lab_mission(tmp_path)
    ms.set_step(m["id"], 1, "done")
    step = ms.get(m["id"])["steps"][0]
    assert step["status"] == "done" and step["evidence"] == "you said so"


def test_finishing_is_refused_while_the_document_shows_gaps(tmp_path):
    m, report = _lab_mission(tmp_path, results_words=40)
    with pytest.raises(ms.MissionError, match="Not finished"):
        ms.finish(m["id"], "done")
    _docx(report, [("Aim", 40), ("Method", 90), ("Results", 160),
                   ("Discussion", 170), ("Conclusion", 90)])
    assert ms.finish(m["id"], "done").startswith("Mission done")
    assert ms.active() is None


def test_a_mission_needs_its_steps(tmp_path):
    notes = tmp_path / "notes.txt"
    notes.write_text("just some notes", encoding="utf-8")
    with pytest.raises(ms.MissionError, match="needs its steps"):
        ms.start("Something", [], files=[str(notes)])
    assert ms.active() is None


def test_a_mission_finishes_itself_when_every_step_is_done(tmp_path):
    m, report = _lab_mission(tmp_path, results_words=40)
    assert not ms.complete_if_done(m["id"])
    _docx(report, [("Aim", 40), ("Method", 90), ("Results", 160), ("Discussion", 170), ("Conclusion", 90)])
    ms.evaluate(m["id"])
    ms.set_step(m["id"], "Read the brief", "done")
    assert not ms.complete_if_done(m["id"])                 # proofread/submit still to do
    ms.set_step(m["id"], "Proofread and submit", "done")
    assert ms.complete_if_done(m["id"])
    assert ms.get(m["id"])["status"] == "done" and ms.active() is None


def test_one_active_mission_at_a_time(tmp_path):
    _lab_mission(tmp_path)
    with pytest.raises(ms.MissionError, match="Already working"):
        ms.start("Another thing", ["step"])


def test_a_missing_file_stops_the_mission_before_anything_is_saved(tmp_path):
    with pytest.raises(ms.MissionError, match="doesn't exist"):
        ms.start("Essay", ["Write intro"], files=[str(tmp_path / "nope.docx")])
    assert ms.active() is None


def test_what_changed_since_the_user_last_looked(tmp_path):
    m, report = _lab_mission(tmp_path, results_words=40)
    _docx(report, [("Aim", 40), ("Method", 90), ("Results", 160), ("Discussion", 30)])
    ms.evaluate(m["id"])
    changes = ms.changes_since_seen(m["id"])
    assert "Results: 40 → 160 words" in changes
    assert "Discussion added (30 words)" in changes
    ms.mark_seen(m["id"])
    assert ms.changes_since_seen(m["id"]) == []


def test_a_mission_survives_a_restart(tmp_path, monkeypatch):
    m, _ = _lab_mission(tmp_path, results_words=40)
    monkeypatch.setattr(ms, "_conn", None)          # a fresh process
    again = ms.active()
    assert again["id"] == m["id"] and again["deadline"] == "tonight"
    assert [s["title"] for s in again["steps"]] == [s["title"] for s in m["steps"]]


def test_the_mission_tool_answers_with_the_true_state(tmp_path):
    from brain.core_runtime import _execute_mission as tool

    report = _docx(tmp_path / "lab3_report.docx", [("Aim", 40), ("Results", 20)])
    brief = _brief(tmp_path / "lab3_brief.txt")
    started = tool({"action": "start", "goal": "Finish lab report",
                    "steps": ["Write the Results", "Write the Conclusion", "Submit it"],
                    "files": [str(report)], "brief": str(brief), "deadline": "Friday 9am"})
    assert started["status"] == "success" and "20/150 words in Results" in started["result"]
    refused = tool({"action": "step", "step": "Results", "status": "done"})
    assert "Not marked done" in refused["result"]
    early = tool({"action": "finish", "status": "done"})
    assert early["status"] == "error" and "Not finished" in early["error"]
    assert tool({"action": "finish", "status": "dropped"})["status"] == "success"
    assert tool({"action": "step", "step": "1"})["status"] == "error"   # nothing active


def test_the_model_is_told_the_mission_in_full_once_then_briefly(tmp_path):
    from brain.core_runtime import CoreRuntime

    _lab_mission(tmp_path, results_words=40)
    rt = CoreRuntime.__new__(CoreRuntime)
    first = rt._mission_context()
    second = rt._mission_context()
    assert first.startswith("Active mission:") and "40/150" in first
    assert second.startswith("Active mission unchanged") and len(second) < 160


def test_files_the_user_names_are_described_to_the_model(tmp_path):
    """Facts to think with: the brief in its own words, the report's
    sections with the words under each. Nothing is decided for the model."""
    from brain.core_runtime import _files_mentioned

    report = _docx(tmp_path / "report.docx", [("Aim", 40), ("Results", 12)])
    brief = _brief(tmp_path / "brief.txt")
    facts = _files_mentioned(f"help me finish this. brief: {brief} report: {report}")
    assert "Discussion: 150-200 words" in facts
    assert "Sections: Aim (40 words), Results (12 words)" in facts
    assert _files_mentioned("what's 12 times 7?") == ""
    ms.start("Finish it", ["Write the Results"], files=[str(report)], brief=[str(brief)])
    assert _files_mentioned(f"how is {report} looking?") == ""    # the mission already says


def test_writing_text_into_a_word_file_is_refused_and_the_file_survives(tmp_path):
    from tools.filesystem.file_manager import FileManager

    report = _docx(tmp_path / "report.docx", [("Aim", 40)])
    before = report.read_bytes()
    with pytest.raises(ValueError, match="would break it"):
        FileManager().write_file(str(report), "Results: g = 9.8")
    with pytest.raises(ValueError, match="would break it"):
        FileManager().append_file(str(report), "more")
    assert report.read_bytes() == before
    assert checks.outline(str(report))["sections"] == {"Aim": 40}


def test_the_brief_for_the_next_sections_travels_with_the_mission(tmp_path):
    m, _ = _lab_mission(tmp_path, results_words=10)
    ctx = ms.context_line(m)
    assert "Brief for Results: Results (about 150 words)" in ctx


def test_reading_a_word_file_gives_its_text_not_its_bytes(tmp_path):
    """read_file on a .docx dumped 12,000 characters of zip bytes into the
    prompt; the next model call ran for minutes."""
    from tools.filesystem.file_manager import FileManager

    report = _docx(tmp_path / "report.docx", [("Aim", 12)])
    text = FileManager().read_file(str(report))
    assert "Aim" in text and "PK" not in text[:10]
    blob = tmp_path / "data.bin"
    blob.write_bytes(b"\x00\x01binary\x00" * 50)
    with pytest.raises(ValueError, match="binary"):
        FileManager().read_file(str(blob))


def test_what_they_wrote_last_travels_with_the_mission(tmp_path):
    from docx import Document

    m, report = _lab_mission(tmp_path, results_words=0)
    d = Document(str(report))
    d.add_paragraph("We measured g = 9.80 m/s2 from the gradient. " * 20)
    d.save(str(report))
    mission = ms.evaluate(m["id"])["mission"]
    assert "What they wrote in Results: “We measured g = 9.80" in ms.context_line(mission)


def test_the_excerpt_keeps_the_students_numbers(tmp_path):
    text = ("We timed the pendulum carefully at each length. " * 6
            + "The measured value of g was 9.80 metres per second squared. "
            + "Repeat timings agreed well and the graph was a straight line through the origin. " * 4)
    excerpt = checks._head_and_tail(text)
    assert "9.80" in excerpt and len(excerpt) <= checks.EXCERPT_CHARS + 20


def test_status_line_and_model_context_are_true_and_short(tmp_path):
    m, _ = _lab_mission(tmp_path, results_words=40)
    line = ms.status_text(m)
    done, total = ms.progress(m)
    assert f"{done} of {total}" in line and "next: Read the brief" in line
    ctx = ms.context_line(m)
    assert "40/150 words in Results" in ctx and "due tonight" in ctx
    assert len(ctx) < 1400


def test_a_length_in_the_models_own_step_is_the_target(tmp_path):
    report = _docx(tmp_path / "report.docx", [("Discussion", 120)])
    m = ms.start("Finish report", ["Write Discussion section (150-200 words)"], files=[str(report)])
    step = m["steps"][0]
    assert step["target"] == 150 and step["status"] == "todo" and "120/150" in step["evidence"]


def test_a_word_brief_listed_with_the_files_is_kept_once_as_the_brief(tmp_path):
    from pathlib import Path
    report = _docx(tmp_path / "report.docx", [("Aim", 40)])
    brief = _docx(tmp_path / "task sheet.docx", [("Results (about 150 words)", 0)])
    m = ms.start("Finish report", ["Write the Results"], files=[str(brief), str(report)], brief=[str(brief)])
    assert sorted((Path(f["path"]).name, f["role"]) for f in m["files"]) == [
        ("report.docx", "work"), ("task sheet.docx", "brief")]
